from __future__ import annotations

from io import BytesIO
import unittest

from openpyxl import load_workbook
import pandas as pd

from src.export.excel_exporter import (
    export_report_table_to_excel,
)
from src.reporter.multi_metrics import (
    build_multi_metric_table,
    validate_calculations_for_question,
)
from src.reporter.table_renderer import _format_cell


class MultiMetricLayoutTests(unittest.TestCase):
    def test_rm_generic_percentage_is_consolidated(
        self,
    ) -> None:
        valid, warnings = validate_calculations_for_question(
            "__rm__",
            [
                "%",
                "RM % Respondentes",
                "RM % Menciones",
            ],
        )

        self.assertEqual(
            valid,
            ["RM % Respondentes", "RM % Menciones"],
        )
        self.assertTrue(
            any("se consolidó" in warning for warning in warnings)
        )

    def test_scalar_metrics_share_one_value_column(self) -> None:
        scale = pd.DataFrame(
            [
                {
                    "banner": "Total",
                    "media": 4.2,
                    "desviacion_estandar": 0.8,
                    "top2box": 0.75,
                    "bottombox": 0.10,
                }
            ]
        )

        result = build_multi_metric_table(
            "Q1",
            calculations=[
                "Media",
                "Desviación estándar",
                "Top2Box",
                "BottomBox",
            ],
            summaries={"scale": scale},
        )

        self.assertEqual(
            result.columns.tolist(),
            ["Respuesta", "Total | Valor"],
        )
        self.assertEqual(
            result["Respuesta"].tolist(),
            [
                "Media",
                "Desviación estándar",
                "Top2Box",
                "BottomBox",
            ],
        )
        self.assertEqual(
            result["Total | Valor"].tolist(),
            [4.2, 0.8, 75.0, 10.0],
        )
        self.assertEqual(
            _format_cell(75.0, "Total | Valor", "Top2Box"),
            "75.0%",
        )

    def test_excel_formats_shared_values_by_metric_row(self) -> None:
        output = BytesIO()
        table = pd.DataFrame(
            {
                "Respuesta": ["Media", "Top2Box"],
                "Total | Valor": [4.2, 75.0],
            }
        )
        export_report_table_to_excel(
            table,
            {"Pregunta": "Q1"},
            output,
        )
        output.seek(0)
        workbook = load_workbook(output, data_only=True)
        worksheet = workbook["Tabla"]

        self.assertEqual(
            worksheet["A1"].value,
            "EXPLORA · REPORTE EJECUTIVO",
        )
        self.assertEqual(worksheet["B7"].value, 4.2)
        self.assertEqual(worksheet["B7"].number_format, "0.00")
        self.assertEqual(worksheet["B8"].value, 0.75)
        self.assertEqual(worksheet["B8"].number_format, "0.0%")
        self.assertEqual(worksheet.max_row, 8)
        self.assertFalse(worksheet["B100"].has_style)


if __name__ == "__main__":
    unittest.main()
