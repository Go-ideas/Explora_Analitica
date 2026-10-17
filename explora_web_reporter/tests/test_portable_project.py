from io import BytesIO
import json
from pathlib import Path
import sqlite3
from zipfile import ZipFile

import numpy as np
import pandas as pd
import pyreadstat
import pytest
from openpyxl import load_workbook

from src.builder.analytic_db_builder import build_analytic_database
from src.export.datamap_exporter import save_final_datamap
from src.export.project_archive import create_project_archive, restore_project_archive
from src.export.project_archive import encode, decode
from src.export.report_exporter import report_to_excel_bytes
from src.readers.datamap_reader import read_datamap, get_review_source_sheet, normalize_review_datamap
from src.readers.spss_reader import read_spss
from src.reporter.filters import FilterChoice
from src.reporter.readiness import analysis_readiness
from src.reporter.report_store import save_report_snapshot
from src.reporter.tabulator import generate_report
from src.utils.text_utils import truthy


def project(folder, repeated):
    folder.mkdir()
    people = np.repeat(np.arange(40), 2) if repeated else np.arange(80)
    frame = pd.DataFrame({'case_key': people.astype(str), 'rating': (np.arange(80) % 9 + 1).astype(float),
                          'profile': people % 2 + 1, 'segment': people % 2 + 1,
                          'benefit_alpha': np.where(np.arange(80) % 3 == 0, 12, 0),
                          'benefit_beta': np.where(np.arange(80) % 4 == 0, 37, 0),
                          'weight': people % 3 + 1.0})
    frame.loc[3, 'rating'] = np.nan
    source = folder / 'responses.sav'
    pyreadstat.write_sav(frame, str(source), variable_value_labels={
        'rating': {i: str(i) for i in range(1, 10)}, 'profile': {1: 'Uno', 2: 'Dos'},
        'segment': {1: 'Uno', 2: 'Dos'}, 'benefit_alpha': {12: 'Alpha'}, 'benefit_beta': {37: 'Beta'}})
    df, meta, _ = read_spss(source)
    rows = [
        {'variable': 'rating', 'pregunta_id': 'SCORE', 'tipo_pregunta': 'Escala', 'escala_min': 1, 'escala_max': 9},
        {'variable': 'profile', 'pregunta_id': 'PROFILE', 'tipo_pregunta': 'RU', 'unidad_analisis': 'persona'},
        {'variable': 'segment', 'pregunta_id': 'SEGMENT', 'tipo_pregunta': 'RU', 'es_banner': True, 'es_filtro': True},
        {'variable': 'benefit_alpha', 'pregunta_id': 'BENEFITS', 'tipo_pregunta': 'RM', 'formato_rm': 'codigo_por_opcion', 'codigo_opcion_rm': 12},
        {'variable': 'benefit_beta', 'pregunta_id': 'BENEFITS', 'tipo_pregunta': 'RM', 'formato_rm': 'codigo_por_opcion', 'codigo_opcion_rm': 37},
        {'variable': 'weight', 'clasificacion_analitica': 'Ponderador', 'es_ponderador': True},
    ]
    for row in rows:
        row.setdefault('unidad_analisis', 'evaluacion')
        row.update(variable_id_persona='case_key', usar_en_dashboard=True)
    raw = pd.DataFrame(rows)
    datamap = folder / 'map.xlsx'
    with pd.ExcelWriter(datamap) as writer:
        raw.to_excel(writer, sheet_name='08_Clasificacion_Analitica', index=False)
    sheets = read_datamap(datamap)
    _, raw = get_review_source_sheet(sheets)
    review = normalize_review_datamap(raw, list(df))
    final = folder / 'final.xlsx'
    save_final_datamap(final, sheets, review, df_spss=df, meta_spss=meta)
    db = folder / 'project.db'
    build_analytic_database(df, meta, review, sheets, db_path=db, export_revision=False)
    return db, source, final, df


@pytest.mark.parametrize('repeated', [False, True])
def test_end_to_end_sources_reference_export_and_recovery(tmp_path, repeated):
    db, source, datamap, df = project(tmp_path / 'original', repeated)
    report = generate_report(db, 'SCORE', calculations=['Media', 'Top2Box'], filters={'segment': [1]})
    expected = df.loc[df.segment.eq(1), 'rating'].dropna()
    assert report.base == len(expected)
    scale = report.summaries['scale'].query("banner == 'Total'").iloc[0]
    assert scale['media'] == pytest.approx(expected.mean())
    rm = generate_report(db, 'BENEFITS', calculations=['RM % Respondentes', 'RM % Menciones'])
    total = rm.summaries['rm'].query("banner == 'Total'")
    assert rm.base == len(df)
    assert total.menciones.sum() == int(df[['benefit_alpha', 'benefit_beta']].ne(0).sum().sum())
    profile = generate_report(db, 'PROFILE', calculations=['n', '%'])
    assert profile.base == (40 if repeated else 80)
    assert load_workbook(BytesIO(report_to_excel_bytes(report))).sheetnames
    saved = save_report_snapshot([], report, 'Escala filtrada')
    state = {'report_question': 'SCORE', 'report_filter_variables': ['segment'],
             'report_filter_values_segment': [FilterChoice(1, 'Uno')],
             'report_calculations_SCORE_v2': ['Media', 'Top2Box'], 'saved_reports': saved,
             'datamap_review_saved': True, 'GH_TOKEN': 'must-not-export'}
    archive = create_project_archive(db, state, source=source, datamap=datamap)
    assert b'must-not-export' not in archive
    restored = restore_project_archive(archive, tmp_path / 'restored')
    assert restored['controls']['report_filter_values_segment'] == [FilterChoice(1, 'Uno')]
    assert restored['datamap_review_saved']
    pd.testing.assert_frame_equal(restored['saved_reports'][0]['report'].table, report.table, check_dtype=False)
    again = generate_report(restored['db_path'], 'SCORE', calculations=['Media', 'Top2Box'], filters={'segment': [1]})
    pd.testing.assert_frame_equal(again.table, report.table)
    assert read_spss(restored['spss_path'])[0].shape == df.shape
    assert read_datamap(restored['datamap_path'])


def test_corruption_and_unexpected_paths_are_rejected(tmp_path):
    db, _, _, _ = project(tmp_path / 'source', False)
    archive = create_project_archive(db, {})
    for extra in ['../outside', 'database.db']:
        buffer = BytesIO(archive)
        with ZipFile(buffer, 'a') as zipped:
            zipped.writestr(extra, b'bad')
        with pytest.raises(ValueError, match='duplicados o no permitidos'):
            restore_project_archive(buffer.getvalue(), tmp_path / 'invalid')
        assert not (tmp_path / 'invalid').exists()
    buffer = BytesIO()
    with ZipFile(BytesIO(archive)) as old, ZipFile(buffer, 'w') as new:
        for name in old.namelist():
            new.writestr(name, b'altered' if name == 'database.db' else old.read(name))
    with pytest.raises(ValueError, match='integridad'):
        restore_project_archive(buffer.getvalue(), tmp_path / 'invalid')


def test_readiness_reports_missing_metadata_and_bad_weights(tmp_path):
    db, _, _, _ = project(tmp_path / 'source', True)
    with sqlite3.connect(db) as conn:
        conn.execute("UPDATE preguntas SET unidad_analisis='' WHERE pregunta_id='SCORE'")
        conn.execute('UPDATE respondentes SET weight=-1 WHERE rowid=1')
    review = analysis_readiness(db).set_index('Revisión')
    assert review.at['Unidad de análisis', 'Estado'] == 'Requiere revisión'
    assert review.at['Ponderación', 'Estado'] == 'Requiere revisión'
    assert '40 personas' in review.at['Identificación de personas', 'Detalle / acción']
    assert review.at['Reglas de aplicación y aprobación metodológica', 'Estado'] == 'Requiere revisión'


@pytest.mark.parametrize('value,expected', [(1.0, True), (0.0, False), (2.0, False), (float('nan'), False), (np.int64(1), True)])
def test_excel_numeric_boolean_flags(value, expected):
    assert truthy(value) is expected


@pytest.mark.parametrize('selected', [False, True])
def test_only_missing_variables_selected_for_use_block_build(monkeypatch, selected):
    from src.ui.page_03_build_database import _validate_datamap_for_build
    import src.ui.page_03_build_database as page
    errors = []
    monkeypatch.setattr(page.st, 'error', errors.append)
    monkeypatch.setattr(page.st, 'warning', lambda message: None)
    datamap = pd.DataFrame([{'variable': 'metadata_only', 'clasificacion_analitica': 'Pregunta analizable',
                            'usar_en_dashboard': selected, 'pregunta_id': 'LOGICAL_PARENT'}])
    assert _validate_datamap_for_build(datamap, pd.DataFrame({'physical': [1]}))['ok'] is (not selected)
    assert bool(errors) is selected


def test_new_upload_clears_previous_analysis_and_canonical_results(monkeypatch):
    import src.ui.page_01_upload as page
    state = {'report_question': 'OLD', 'paired_metric': 'Media', 'canonical_result': object(),
             'canonical_results': {'OLD': object()}, 'current_report': object(),
             'project_download': b'old', 'db_path': 'old.db', 'saved_reports': [],
             'project_navigation': '1. Estado del proyecto'}
    monkeypatch.setattr(page.st, 'session_state', state)
    page._reset_project_state()
    assert state == {'project_navigation': '1. Estado del proyecto'}


def test_interface_keeps_choices_across_pages_and_exports_saved_analysis(tmp_path):
    from streamlit.testing.v1 import AppTest
    from src.reporter.filters import filter_choices
    db, _, _, _ = project(tmp_path / 'source', False)
    app = AppTest.from_file(str(Path(__file__).resolve().parents[1] / 'app.py'))
    app.session_state['db_path'] = str(db)
    app.run(timeout=60)
    app.selectbox(key='report_question').select('SCORE').run(timeout=60)
    app.multiselect(key='report_filter_variables').set_value(['segment']).run(timeout=60)
    choice = next(value for value in filter_choices(db, 'segment') if value.value == 1)
    app.multiselect(key='report_filter_values_segment').set_value([choice]).run(timeout=60)
    app.button(key='generate_report').click().run(timeout=60)
    app.button(key='report_save_report_SCORE').click().run(timeout=60)
    app.radio(key='project_navigation').set_value('5. Exportar').run(timeout=60)
    assert not app.exception, [item.message for item in app.exception]
    assert app.session_state['report_question'] == 'SCORE'
    app.button(key='project_prepare').click().run(timeout=60)
    assert not app.exception and not app.error
    restored = restore_project_archive(app.session_state['project_download'], tmp_path / 'restored')
    assert restored['controls']['report_question'] == 'SCORE'
    assert restored['controls']['report_filter_values_segment'] == [choice]
    assert len(restored['saved_reports']) == 1
    assert restored['saved_reports'][0]['report'].filters == {'segment': [1]}


def test_snapshot_preserves_tiny_probabilities_and_full_float_precision():
    frame = pd.DataFrame({'p': [1e-25, 0.0], 'value': [1.1218655663016717, 7.725]})
    pd.testing.assert_frame_equal(decode(encode(frame)), frame, check_exact=True)
