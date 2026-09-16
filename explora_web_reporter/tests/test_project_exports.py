from __future__ import annotations

from io import BytesIO
from pathlib import Path
import tempfile
import unittest
from zipfile import ZipFile

import pandas as pd
import plotly.graph_objects as go
from openpyxl import load_workbook

from src.export.project_exporter import (
    executive_export_bytes,
    saved_reports_export_bytes,
    technical_export_bytes,
)
from src.export.report_exporter import (
    report_to_excel_bytes,
    reports_to_excel_bytes,
)
from src.reporter.report_store import (
    delete_report_snapshots,
    save_report_snapshot,
)
from src.reporter.tabulator import ReportResult


class ProjectExportTests(unittest.TestCase):
    @staticmethod
    def _report() -> ReportResult:
        return ReportResult(
            question_id="Q1",
            title="Pregunta de prueba",
            section="Perfil",
            question_type="RU",
            calculation="Porcentaje",
            calculations=["n", "%"],
            base=100,
            banner=None,
            banners=[],
            filters={},
            filter_summary="Sin filtros",
            filter_variable=None,
            filter_display=None,
            ponderador=None,
            display_mode="n y %",
            response_order="Orden de Value Labels",
            confidence=None,
            significance_display="none",
            table=pd.DataFrame(
                {
                    "Respuesta": ["Sí"],
                    "Total | n": [100],
                    "Total | %": [100.0],
                }
            ),
            summary=pd.DataFrame(),
            summaries={},
            figure=go.Figure(go.Bar(x=["Sí"], y=[100])),
            notes=["Nota metodológica"],
        )

    def test_executive_export_contains_table_and_chart(
        self,
    ) -> None:
        report = self._report()

        with ZipFile(BytesIO(executive_export_bytes(report))) as archive:
            names = set(archive.namelist())
            chart = archive.read("Q1_Grafico.html")

        self.assertIn("Q1_Analisis_Explora.xlsx", names)
        self.assertIn("Q1_Grafico.html", names)
        self.assertIn("LEEME.txt", names)
        self.assertIn(b"plotly", chart.lower())

    def test_technical_export_preserves_complete_files(
        self,
    ) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            source = Path(temp_dir) / "Datamap_Final.xlsx"
            source.write_bytes(b"contenido")
            payload = technical_export_bytes(
                [("Datamap_Final.xlsx", source)]
            )

        with ZipFile(BytesIO(payload)) as archive:
            self.assertEqual(
                archive.read("Datamap_Final.xlsx"),
                b"contenido",
            )
            self.assertIn("CONTENIDO.txt", archive.namelist())

    def test_saved_reports_can_be_exported_and_deleted(
        self,
    ) -> None:
        report = self._report()
        saved = save_report_snapshot(
            [], report, "Primer análisis"
        )
        saved = save_report_snapshot(
            saved, report, "Segundo análisis"
        )
        first_id = saved[0]["id"]

        with ZipFile(
            BytesIO(saved_reports_export_bytes(saved))
        ) as archive:
            names = archive.namelist()

        self.assertIn("Todas_las_tablas.xlsx", names)
        self.assertTrue(
            any(
                name.endswith("_Tabla.xlsx")
                and "Primer_análisis" in name
                for name in names
            )
        )
        self.assertTrue(
            any(
                name.endswith("_Grafico.html")
                and "Segundo_análisis" in name
                for name in names
            )
        )

        with ZipFile(
            BytesIO(saved_reports_export_bytes(saved))
        ) as archive:
            combined_book = load_workbook(
                BytesIO(
                    archive.read("Todas_las_tablas.xlsx")
                ),
                data_only=True,
            )
        individual_book = load_workbook(
            BytesIO(report_to_excel_bytes(report)),
            data_only=True,
        )
        combined_sheet = combined_book[
            combined_book.sheetnames[0]
        ]
        individual_sheet = individual_book["Tabla"]
        self.assertEqual(
            combined_sheet["A1"].value,
            individual_sheet["A1"].value,
        )
        self.assertEqual(
            [
                combined_sheet.cell(5, column).value
                for column in range(1, 4)
            ],
            [
                individual_sheet.cell(5, column).value
                for column in range(1, 4)
            ],
        )
        self.assertFalse(combined_sheet["B100"].has_style)

        remaining = delete_report_snapshots(
            saved, [first_id]
        )
        self.assertEqual(len(remaining), 1)
        self.assertEqual(
            remaining[0]["name"], "Segundo análisis"
        )

    def test_integrated_significance_is_not_duplicated(
        self,
    ) -> None:
        report = self._report()
        report.significance_display = "integrated"
        report.table["Total | Sig."] = ["A"]
        report.significance = pd.DataFrame(
            {"Respuesta": ["Sí"], "Total": ["A"]}
        )
        report.significance_legend = pd.DataFrame(
            {"Letra": ["A"], "Columna": ["Total"]}
        )

        with pd.ExcelFile(
            BytesIO(reports_to_excel_bytes([report]))
        ) as workbook:
            frame = pd.read_excel(
                workbook,
                sheet_name=workbook.sheet_names[0],
                header=None,
            )
        values = {
            str(value)
            for value in frame.to_numpy().ravel()
            if pd.notna(value)
        }

        self.assertNotIn("Significancia", values)
        self.assertIn("EXPLORA · REPORTE EJECUTIVO", values)
        self.assertIn("Sig.", values)

        report.significance_display = "separate"
        with pd.ExcelFile(
            BytesIO(reports_to_excel_bytes([report]))
        ) as workbook:
            sheet_names = workbook.sheet_names
            significance_sheet = pd.read_excel(
                workbook,
                sheet_name=sheet_names[1],
                header=None,
            )
        significance_values = {
            str(value)
            for value in significance_sheet.to_numpy().ravel()
            if pd.notna(value)
        }
        self.assertEqual(len(sheet_names), 2)
        self.assertTrue(
            any(
                "Diferencias significativas" in value
                for value in significance_values
            )
        )


if __name__ == "__main__":
    unittest.main()
