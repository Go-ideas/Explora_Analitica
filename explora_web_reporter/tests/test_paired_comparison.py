from io import BytesIO
from types import SimpleNamespace

import numpy as np
import pandas as pd
import pytest
from scipy.stats import binomtest, ttest_rel
from openpyxl import load_workbook

from src.builder.analytic_db_builder import build_analytic_database
from src.database.db_reader import read_table
from src.readers.datamap_reader import normalize_review_datamap
from src.reporter.paired_comparison import PairedCondition, compare_paired


def build(tmp_path, *, duplicate=False):
    rows=[]
    for person in range(40):
        for group in [1,2]:
            value=3+(person%5) if group==1 else 2+((person+1)%5)
            rows.append({'case_key':f'P{person}', 'arm':group,'score':float(value),
                         'followup':float(value+1), 'select':12 if person%4<group else 0})
    df=pd.DataFrame(rows).sample(frac=1,random_state=17).reset_index(drop=True)
    if duplicate:df=pd.concat([df,df.iloc[[0]]],ignore_index=True)
    raw=pd.DataFrame([
        {'variable':'score','pregunta_id':'SAT','tipo_pregunta':'Escala','id_respondente':'case_key',
         'tipo_eje_fila':'Visita','eje_fila_grid':'Antes','usar_en_dashboard':True},
        {'variable':'followup','pregunta_id':'SAT','tipo_pregunta':'Escala','id_respondente':'case_key',
         'tipo_eje_fila':'Visita','eje_fila_grid':'Después','usar_en_dashboard':True},
        {'variable':'arm','tipo_pregunta':'RU','es_banner':True,'es_filtro':True,'usar_en_dashboard':True},
        {'variable':'select','pregunta_id':'BEN','tipo_pregunta':'RM','formato_rm':'codigo_por_opcion',
         'codigo_opcion_rm':12,'usar_en_dashboard':True},
    ])
    dm=normalize_review_datamap(raw,list(df))
    meta=SimpleNamespace(variable_value_labels={
        'score':{i:str(i) for i in range(1,10)},'followup':{i:str(i) for i in range(1,10)},
        'select':{12:'Benefit'}},column_names_to_labels={})
    db=tmp_path/'paired.db'
    build_analytic_database(df,meta,dm,db_path=db,export_revision=False)
    return db,df


def sides(question='SAT::score'):
    return PairedCondition(question,{'arm':[1]}),PairedCondition(question,{'arm':[2]})


def test_means_pair_by_identity_not_row_order(tmp_path):
    db,df=build(tmp_path)
    r=compare_paired(db,*sides())
    expected=df.pivot(index='case_key',columns='arm',values='score')
    row=r.table.iloc[0]
    assert row['Pares completos']==40
    assert row['p bilateral']==pytest.approx(ttest_rel(expected[1],expected[2]).pvalue)
    assert row['Diferencia A−B']==pytest.approx((expected[1]-expected[2]).mean())
    ci=ttest_rel(expected[1],expected[2]).confidence_interval(confidence_level=.95)
    assert row['IC inferior']==pytest.approx(ci.low)
    assert row['IC superior']==pytest.approx(ci.high)
    assert r.config['Método']=='t pareada bilateral'


def test_exact_mcnemar_uses_discordant_person_pairs(tmp_path):
    db,df=build(tmp_path)
    r=compare_paired(db,*sides('BEN'),metric='Selección RM',option_code=12)
    row=r.table.iloc[0]
    assert row['Pares completos']==40
    assert row['A sí/B no']==0 and row['A no/B sí']==10
    assert row['p bilateral']==pytest.approx(binomtest(0,10,.5).pvalue)
    assert row['Valor A (pares)']==.25 and row['Valor B (pares)']==.5


def test_binary_boxes_compare_only_matched_people(tmp_path):
    db,df=build(tmp_path)
    r=compare_paired(db,*sides(),metric='Top2Box')
    assert r.table.iloc[0]['p bilateral']==1.0 # Neither arm has codes 8 or 9.


def test_repeated_times_use_declared_parent_and_export_method(tmp_path):
    db,_=build(tmp_path)
    r=compare_paired(db,PairedCondition('SAT::score',{'arm':[1]}),PairedCondition('SAT::followup',{'arm':[1]}))
    assert r.table.iloc[0]['Diferencia A−B']==-1
    assert r.table.iloc[0]['p bilateral']==0
    wb=load_workbook(BytesIO(r.excel_bytes()),data_only=True)
    assert 'Tabla' in wb.sheetnames and 'Configuracion' in wb.sheetnames
    assert any('t pareada' in str(cell.value) for row in wb['Configuracion'] for cell in row)
    assert any('pares' in str(cell.value) for row in wb['Notas_Metodologicas'] for cell in row)


def test_duplicate_observations_require_explicit_person_mean(tmp_path):
    db,df=build(tmp_path,duplicate=True)
    with pytest.raises(ValueError,match='varias evaluaciones'):
        compare_paired(db,*sides())
    import sqlite3
    people=read_table(db,'respondentes')
    extra=people.loc[people.row_index.eq(len(df)-1),'id_respondente'].iloc[0]
    df.loc[len(df)-1,'score']-=1
    with sqlite3.connect(db) as conn:
        conn.execute('UPDATE escalas_long SET valor=valor-1 WHERE pregunta_id=? AND id_respondente=?',('SAT::score',extra))
    expected=df.groupby(['case_key','arm']).score.mean().unstack()
    result=compare_paired(db,*sides(),duplicate_policy='person_mean')
    assert result.table.iloc[0]['Pares completos']==40
    assert result.table.iloc[0]['Diferencia A−B']==pytest.approx((expected[1]-expected[2]).mean())
    assert result.table.iloc[0]['p bilateral']==pytest.approx(ttest_rel(expected[1],expected[2]).pvalue)
    with pytest.raises(ValueError,match='una observación'):
        compare_paired(db,*sides(),metric='Top2Box',duplicate_policy='person_mean')


def test_missing_pairs_do_not_become_independent_samples(tmp_path):
    db,_=build(tmp_path)
    # There are no people present in both these mutually exclusive identity filters.
    r=compare_paired(db,PairedCondition('SAT::score',{'id_persona':['P0','P1']}),
                     PairedCondition('SAT::score',{'id_persona':['P2','P3']}),duplicate_policy='person_mean')
    assert r.table.iloc[0]['Pares completos']==0
    assert r.table.iloc[0]['Resultado']=='Base insuficiente'
    assert r.table.iloc[0]['p bilateral'] is None


@pytest.mark.parametrize('kwargs',[{'min_pairs':1},{'confidence':1},{'metric':'NPS'},
                                  {'duplicate_policy':'first'},{'metric':'Selección RM'}])
def test_invalid_requests_fail_closed(tmp_path,kwargs):
    db,_=build(tmp_path)
    with pytest.raises(ValueError):compare_paired(db,*sides(),**kwargs)


def test_insufficient_pairs_are_not_reported_as_nonsignificant(tmp_path):
    db,_=build(tmp_path)
    r=compare_paired(db,*sides(),min_pairs=41)
    assert r.table.iloc[0]['Resultado']=='Base insuficiente'
    assert r.table.iloc[0]['p bilateral'] is None


def test_overlap_and_missing_filters_are_not_silently_ignored(tmp_path):
    db,_=build(tmp_path)
    for left,right in [(PairedCondition('SAT::score'),sides()[1]),
                        (PairedCondition('SAT::score',{'not_a_column':[1]}),sides()[1])]:
        with pytest.raises(ValueError):compare_paired(db,left,right,duplicate_policy='person_mean')


@pytest.mark.parametrize('question,metric,code',[('BEN','Media',None),('SAT::score','Selección RM',12)])
def test_metric_must_match_both_question_types(tmp_path,question,metric,code):
    db,_=build(tmp_path)
    with pytest.raises(ValueError,match='métrica'):
        compare_paired(db,*sides(question),metric=metric,option_code=code)


def test_missing_response_reduces_pair_base_and_keeps_counts(tmp_path):
    import sqlite3
    db,df=build(tmp_path)
    people=read_table(db,'respondentes')
    absent=people.loc[(people.id_persona=='P0') & (people.arm==2),'id_respondente'].iloc[0]
    with sqlite3.connect(db) as conn:
        conn.execute('DELETE FROM escalas_long WHERE pregunta_id=? AND id_respondente=?',('SAT::score',absent))
    r=compare_paired(db,*sides())
    row=r.table.iloc[0]
    expected=df.pivot(index='case_key',columns='arm',values='score').drop(index='P0')
    assert row['Pares completos']==39 and row['Sin pareja A']==1 and row['Sin pareja B']==0
    assert row['p bilateral']==pytest.approx(ttest_rel(expected[1],expected[2]).pvalue)


@pytest.mark.parametrize('mutation', ['labels','bounds','parent','excluded'])
def test_incompatible_definitions_or_exclusions_require_validation(tmp_path,mutation):
    import sqlite3
    db,_=build(tmp_path)
    with sqlite3.connect(db) as conn:
        if mutation=='labels':conn.execute("UPDATE opciones SET label='reversed' WHERE pregunta_id=? AND codigo=1",('SAT::followup',))
        elif mutation=='bounds':conn.execute('UPDATE escalas_long SET escala_max=10 WHERE pregunta_id=?',('SAT::followup',))
        elif mutation=='parent':conn.execute("UPDATE preguntas SET pregunta_padre='OTHER' WHERE pregunta_id=?",('SAT::followup',))
        else:conn.execute('UPDATE preguntas SET usar_en_dashboard=0 WHERE pregunta_id=?',('SAT::followup',))
    with pytest.raises(ValueError):
        compare_paired(db,PairedCondition('SAT::score',{'arm':[1]}),PairedCondition('SAT::followup',{'arm':[1]}))


def test_generated_row_identifiers_do_not_prove_person_identity(tmp_path):
    import sqlite3
    db,_=build(tmp_path)
    with sqlite3.connect(db) as conn:conn.execute("UPDATE respondentes SET id_original='' WHERE id_persona='P0'")
    with pytest.raises(ValueError,match='identificador de persona válido'):
        compare_paired(db,*sides())


def test_rm_without_observed_universe_is_not_given_a_test(tmp_path):
    import sqlite3
    db,_=build(tmp_path)
    with sqlite3.connect(db) as conn:conn.execute('DELETE FROM bases_multirrespuesta')
    with pytest.raises(ValueError,match='base observada explícita'):
        compare_paired(db,*sides('BEN'),metric='Selección RM',option_code=12)


def test_nan_is_not_a_declared_rm_option(tmp_path):
    db,_=build(tmp_path)
    with pytest.raises(ValueError,match='código RM'):
        compare_paired(db,*sides('BEN'),metric='Selección RM',option_code=float('nan'))


def test_nonfinite_response_does_not_become_a_nonsignificant_result(tmp_path):
    import sqlite3
    db,_=build(tmp_path)
    with sqlite3.connect(db) as conn:conn.execute('UPDATE escalas_long SET valor=? WHERE pregunta_id=?',(float('inf'),'SAT::score'))
    with pytest.raises(ValueError,match='no finitos'):
        compare_paired(db,*sides())
