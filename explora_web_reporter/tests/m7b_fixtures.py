"""Independent, preallocated generic test Master construction, never production rendering."""
from dataclasses import replace
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet.table import Table
from openpyxl.workbook.defined_name import DefinedName

from src.analytics_core.result_identity import canonical_payload, request_identity, result_fingerprint, stable_json
from src.contracts.models import ReleaseMetadata
from src.contracts.vocabulary import ReleaseLifecycle, ReleaseMode
from src.excel_renderer.renderer import (
    BUILD, MASTER, NUMERIC, ROLES, VISUAL, MasterManifest, RenderRequest, Slot,
    TableBinding, audit_columns, configuration_fingerprint, sha, sha_bytes,
)
from m6_fixtures import canonical_result

FIXTURE = Path(__file__).parent / "fixtures" / "m7b" / "generic_macro_renderer_fixture.xlsm"
FIXTURE_SHA = "846fc71ae54e02a99a8278f662dcc15584c1e4575030398ddb4825363471bc06"
VBA_SHA = "0f879b60ed12315085bc722c3f59163f86ce24609e3ba6ad379069c44779b758"


def release(checksum):
    return ReleaseMetadata(ReleaseLifecycle.RELEASED, ReleaseMode.MANUAL,
                           "GENERIC_TEST_AUTHORITY", "2026-09-18", "M7_TEST_ONLY", "V1", checksum)


def source_result(**kwargs):
    run = canonical_result(**kwargs)
    _, request_fp = request_identity(canonical_payload(run.request))
    run = replace(run, request=replace(run.request, request_fingerprint=request_fp))
    fp = result_fingerprint(run)
    return replace(run, result_fingerprint=fp, manifest=replace(run.manifest, result_fingerprint=fp))


def build_test_master(path, slots=None):
    # Whole-workbook construction is test infrastructure ONLY. Renderer never saves through openpyxl.
    w = load_workbook(FIXTURE, keep_vba=True)
    roles = []
    for role in ROLES:
        if role in ("presentation", "configuration", "navigation"):
            table = next(iter(w[role].tables.values()))
            roles.append(TableBinding(role, role, table.name, 11 if role == "configuration" else 1))
            continue
        s = w[role]
        for name in list(s.tables):
            del s.tables[name]
        columns = audit_columns(role)
        for col, name in enumerate(columns, 1):
            s.cell(1, col, name)
        capacity = 8
        for row in range(2, capacity + 2):
            for col in range(1, len(columns)+1):
                s.cell(row, col, "STALE_GENERIC_ONLY")
        from openpyxl.utils import get_column_letter
        name = "Generic_" + role
        s.add_table(Table(displayName=name, ref=f"A1:{get_column_letter(len(columns))}{capacity+1}"))
        roles.append(TableBinding(role, role, name, capacity))
    presentation = w["presentation"]
    for cell in ("F1", "G1", "H1", "G2"):
        presentation[cell] = "GENERIC_DECLARED_ONLY"
    w.defined_names.add(DefinedName("ValueStatusSlot", attr_text="'presentation'!$G$1"))
    w.defined_names.add(DefinedName("TextStatusSlot", attr_text="'presentation'!$H$1"))
    w.defined_names.add(DefinedName("OtherStatusSlot", attr_text="'presentation'!$G$2"))
    w.defined_names.add(DefinedName("TechnicalLabelSlot", attr_text="'presentation'!$A$2"))
    registry = w["configuration"]
    for index, role in enumerate(ROLES, 2):
        binding = next(b for b in roles if b.role_id == role)
        registry.cell(index, 5, binding.table)
        registry.cell(index, 6, binding.capacity)
    w.save(path)
    w.close()
    w.vba_archive.close()
    if slots is None:
        slots = (Slot("official_value", "presentation", "numeric", "exact_numeric",
                      name="RendererSlot", status_name="ValueStatusSlot"),)
    slots = tuple(replace(s, label_name="TechnicalLabelSlot") if s.kind == "numeric" and s.label_name is None
                  else s for s in slots)
    return MasterManifest("GENERIC_TEST_MASTER", "V1", sha(path), MASTER, (VISUAL,), tuple(roles),
                          tuple(slots), (("presentation", "E1", "=1+1"),))


def request_for(master, run=None, bindings=None, profiles=None, significance=()):
    run = run or source_result()
    if bindings is None:
        bindings = [{"result_run_id": run.result_run_id, "record_role": "value", "record_id": run.values[0].value_id,
                     "field_name": "estimate", "slot_id": master.writable_slots[0].slot_id}]
    profiles = profiles or ["PROPORTION"] * len(bindings)
    refs = [sha_bytes(raw) for raw in significance]
    spec = {"schema_version": VISUAL, "visual_spec_id": "GENERIC_VISUAL", "revision": 1,
            "project_id": run.project_id,
            "result_refs": [{"result_run_id": run.result_run_id, "result_fingerprint": run.result_fingerprint,
                             "request_fingerprint": run.request.request_fingerprint}],
            "master_interface_version": MASTER, "numeric_profile_version": NUMERIC, "provenance_refs": [],
            "sections": [{"section_id": "generic_section", "order": 0, "sheet_role": "presentation",
                          "title": {"text": "GENERIC", "language": "en"},
                          "visuals": [{"visual_id": f"visual_{i}", "order": i, "role": "result_table",
                                       "source_bindings": [b], "labels": [], "display_profile_ref": p,
                                       "warning_refs": [], "significance_presentation_refs": refs}
                                      for i, (b, p) in enumerate(zip(bindings, profiles))]}]}
    request = RenderRequest((run,), stable_json(spec), master, master.writable_slots,
                            "GENERIC_CONFIGURATION", "V1", release("PENDING"), release(master.artifact_sha256),
                            significance_json=significance, renderer_build=BUILD)
    return resign(request)


def resign(request):
    return replace(request, configuration_release=release(configuration_fingerprint(request)))
