from dataclasses import replace
import hashlib
import json
from pathlib import Path
import subprocess
from zipfile import ZipFile

from openpyxl import load_workbook
import pytest

from src.analytics_core.result_identity import result_fingerprint, stable_json
from src.excel_renderer import (
    Bounds, DynamicRegionBinding, DynamicRenderRequest, RenderError,
    lineage_from, load_dynamic_qualified_master, load_qualified_master,
    plan_dynamic_render, plan_render, render, render_dynamic,
)
from src.excel_renderer.dynamic_region import DYNAMIC_BINDING_SCHEMA
from src.excel_renderer.production import significance_qualification_request
from src.excel_renderer.renderer import SIGNIFICANCE, sha, sha_bytes
from src.excel_renderer.qualification import vba_sha256
from test_gate39_significance_presentation import core_result

ROOT = Path(__file__).parents[1]
BASELINE = ROOT / "production_masters/EXPLORA_PRODUCTION_MASTER_V1.xlsm"
BASELINE_MANIFEST = BASELINE.with_suffix(".qualification.json")
CANDIDATE = ROOT / "production_masters/EXPLORA_PRODUCTION_MASTER_V1_2.xlsm"
CANDIDATE_MANIFEST = CANDIDATE.with_suffix(".qualification.json")
DYNAMIC_MANIFEST = CANDIDATE.with_suffix(".dynamic.qualification.json")
BASELINE_SHA = "a8313adf706016b30b8837a7bc42cb80c7ccc281f2663616a7f22b8daa8d10dd"
CANDIDATE_SHA = "f3a11f291b6c661f0c937e9c95d7f227c351d7261b2e01aa44dfd4167d5c869c"
VBA_SHA = "0f879b60ed12315085bc722c3f59163f86ce24609e3ba6ad379069c44779b758"
RESULT_REGION = "PRODUCTION_DYNAMIC_RESULTS_V1"
PROVENANCE_REGION = "PRODUCTION_DYNAMIC_PROVENANCE_V1"


def expanded_core_result(count=8):
    run = core_result()
    source = tuple(v for v in run.values if v.slice_id in
                   (run.comparisons[0].left_slice_id, run.comparisons[0].right_slice_id))
    values = tuple(replace(source[index % len(source)], value_id=f"gate42_value_{index + 1}",
                           semantic_order=index + 1) for index in range(count))
    unsigned = replace(run, values=values)
    fp = result_fingerprint(unsigned)
    return replace(unsigned, result_fingerprint=fp,
                   manifest=replace(unsigned.manifest, result_fingerprint=fp))


def token_envelope(run):
    comparison = run.comparisons[0]
    bindings = [{"comparison_id": comparison.comparison_id, "family_id": comparison.family_id,
        "test_id": comparison.test_id, "test_version": comparison.test_version,
        "value_id": value.value_id, "slice_id": value.slice_id,
        "member_id": comparison.left_member_id if value.slice_id == comparison.left_slice_id
        else comparison.right_member_id, "direction": comparison.direction,
        "token_id": f"gate42_marker_{index + 1}", "display_text": "*"}
        for index, value in enumerate(run.values)]
    return stable_json({"schema_version": SIGNIFICANCE, "token_set_id": "GATE42_PRODUCTION_MARKERS_V1",
        "revision": 1, "authority_ref": "GATE42_CORE_B2", "result_run_id": run.result_run_id,
        "result_fingerprint": run.result_fingerprint, "comparison_bindings": bindings,
        "legend": [{"token_id": item["token_id"], "meaning": "Canonical Core significance marker",
                    "comparison_refs": [comparison.comparison_id]} for item in bindings]})


def dynamic_request(rows, columns, *, run=None, qualification=None, lineage=None):
    run = run or expanded_core_result()
    qualification = qualification or load_dynamic_qualified_master(DYNAMIC_MANIFEST, CANDIDATE)
    ids = tuple(value.value_id for value in run.values[:rows])
    fields = ("estimate", "status", "significance")[:columns]
    profiles = ("PROPORTION", "LITERAL", "SIGNIFICANCE_TOKEN")[:columns]
    significance_json = ()
    refs = ()
    if columns == 3:
        raw = token_envelope(run)
        significance_json, refs = (raw,), (sha_bytes(raw),)
    bindings = (
        DynamicRegionBinding(RESULT_REGION, run.result_run_id, "value", ids, fields,
                             rows, columns, profiles, refs),
        DynamicRegionBinding(PROVENANCE_REGION, run.result_run_id, "provenance",
                             (run.result_run_id,), ("result_run_id",), 1, 1, ("LITERAL",)),
    )
    return DynamicRenderRequest((run,), qualification, DYNAMIC_BINDING_SCHEMA,
                                bindings, significance_json, lineage)


def copy_candidate(tmp_path):
    path = tmp_path / "candidate.xlsm"
    path.write_bytes(CANDIDATE.read_bytes())
    return path


def test_t42_01_02_baseline_identity_candidate_lineage_and_reproducibility():
    assert sha(BASELINE) == BASELINE_SHA
    assert vba_sha256(BASELINE) == VBA_SHA
    assert CANDIDATE != BASELINE and sha(CANDIDATE) == CANDIDATE_SHA
    assert vba_sha256(CANDIDATE) == VBA_SHA
    dynamic = json.loads(DYNAMIC_MANIFEST.read_text())
    assert (dynamic["baseline_version"], dynamic["baseline_sha256"], dynamic["master_version"]) == (
        "1.1.0", BASELINE_SHA, "1.2.0")
    before = CANDIDATE.read_bytes()
    subprocess.run(["python", "scripts/build_gate42_production_master.py"], cwd=ROOT, check=True)
    assert CANDIDATE.read_bytes() == before
    assert sha(BASELINE) == BASELINE_SHA


@pytest.mark.parametrize("rows,columns", [(1, 1), (4, 2), (8, 3)])
def test_t42_03_04_06_minimum_nominal_maximum_physical_readback(tmp_path, rows, columns):
    source = copy_candidate(tmp_path)
    request = dynamic_request(rows, columns, qualification=replace(
        load_dynamic_qualified_master(DYNAMIC_MANIFEST, CANDIDATE), master_sha256=sha(source)))
    # The copied bytes retain the qualified candidate identity.
    request = replace(request, qualification=replace(request.qualification, master_sha256=CANDIDATE_SHA))
    plan = plan_dynamic_render(request, source)
    operation = next(o for o in plan.dynamic_operations
                     if o.region_id == RESULT_REGION and o.operation_type == "VALIDATE_REGION")
    assert (operation.after_bounds, operation.requested_rows, operation.requested_columns) == (
        Bounds(10, 2, 9 + columns, 1 + rows), rows, columns)
    output = tmp_path / f"extent-{rows}-{columns}.xlsm"
    evidence = render_dynamic(request, source, output, plan=plan)
    assert evidence["qa"] == "PASS" and evidence["output_validation_sha256"]
    workbook = load_workbook(output, keep_vba=True)
    try:
        sheet = workbook["presentation"]
        assert sheet.cell(1 + rows, 9 + columns).value is not None
        assert sheet["N2"].value == request.results[0].result_run_id
        if columns == 3:
            assert all(sheet.cell(2 + index, 12).value == "*" for index in range(rows))
    finally:
        workbook.close(); workbook.vba_archive.close()
    assert vba_sha256(output) == VBA_SHA


def test_t42_05_07_08_expansion_contraction_reexpansion_history_independent(tmp_path):
    source = copy_candidate(tmp_path)
    qualification = load_dynamic_qualified_master(DYNAMIC_MANIFEST, CANDIDATE)
    direct_request = dynamic_request(8, 3, qualification=qualification)
    direct = tmp_path / "direct-expanded.xlsm"
    direct_evidence = render_dynamic(direct_request, source, direct)
    small_request = dynamic_request(2, 2, qualification=qualification)
    small = tmp_path / "small.xlsm"
    small_plan = plan_dynamic_render(small_request, source)
    small_evidence = render_dynamic(small_request, source, small, plan=small_plan)
    expanded_request = dynamic_request(8, 3, qualification=qualification,
        lineage=lineage_from(small_plan, small_evidence))
    expanded = tmp_path / "expanded.xlsm"
    expanded_plan = plan_dynamic_render(expanded_request, small)
    expanded_evidence = render_dynamic(expanded_request, small, expanded, plan=expanded_plan)
    contracted_request = dynamic_request(2, 2, qualification=qualification,
        lineage=lineage_from(expanded_plan, expanded_evidence))
    contracted = tmp_path / "contracted.xlsm"
    contracted_plan = plan_dynamic_render(contracted_request, expanded)
    assert any(o.operation_type == "CLEAR_OWNED_STALE" for o in contracted_plan.dynamic_operations)
    contracted_evidence = render_dynamic(contracted_request, expanded, contracted, plan=contracted_plan)
    reexpand_request = dynamic_request(8, 3, qualification=qualification,
        lineage=lineage_from(contracted_plan, contracted_evidence))
    reexpanded = tmp_path / "reexpanded.xlsm"
    reexpanded_evidence = render_dynamic(reexpand_request, contracted, reexpanded)
    assert direct_evidence["logical_output_sha256"] == expanded_evidence["logical_output_sha256"]
    assert direct_evidence["logical_output_sha256"] == reexpanded_evidence["logical_output_sha256"]
    workbook = load_workbook(contracted, keep_vba=True)
    try:
        assert workbook["presentation"]["J4"].value is None
        assert workbook["presentation"]["L2"].value is None
    finally:
        workbook.close(); workbook.vba_archive.close()


def test_t42_09_overflow_and_no_partial_mutation(tmp_path):
    source = copy_candidate(tmp_path)
    before = sha(source)
    output = tmp_path / "overflow.xlsm"
    request = dynamic_request(8, 3)
    bad_binding = replace(request.bindings[0], requested_rows=9,
        ordered_record_ids=(*request.bindings[0].ordered_record_ids, request.bindings[0].ordered_record_ids[0]))
    request = replace(request, bindings=(bad_binding, request.bindings[1]))
    with pytest.raises(RenderError, match="overflow"):
        render_dynamic(request, source, output)
    assert sha(source) == before and not output.exists()


def test_t42_10_collision_and_no_partial_mutation(tmp_path):
    source = copy_candidate(tmp_path)
    qualification = load_dynamic_qualified_master(DYNAMIC_MANIFEST, CANDIDATE)
    declaration = replace(qualification.declarations[0], anchor_name="ProductionSignificanceSlot",
                          owned_envelope=Bounds(5, 2, 7, 9))
    style = replace(qualification.style_policies[0], template_cell="presentation!E2")
    qualification = replace(qualification, declarations=(declaration, qualification.declarations[1]),
                            style_policies=(style, qualification.style_policies[1]))
    before = sha(source)
    output = tmp_path / "collision.xlsm"
    with pytest.raises(RenderError, match="fixed-slot collision"):
        render_dynamic(dynamic_request(2, 2, qualification=qualification), source, output)
    assert sha(source) == before and not output.exists()


def test_t42_11_12_13_14_significance_provenance_readback_determinism(tmp_path):
    source = copy_candidate(tmp_path)
    request = dynamic_request(4, 3)
    first, second = tmp_path / "first.xlsm", tmp_path / "second.xlsm"
    plan_a, plan_b = plan_dynamic_render(request, source), plan_dynamic_render(request, source)
    assert plan_a == plan_b and plan_a.plan_sha256 == plan_b.plan_sha256
    evidence_a = render_dynamic(request, source, first, plan=plan_a)
    evidence_b = render_dynamic(request, source, second, plan=plan_b)
    assert evidence_a["logical_output_sha256"] == evidence_b["logical_output_sha256"]
    assert evidence_a["output_validation_sha256"] == evidence_b["output_validation_sha256"]
    assert any(write.token_id and write.comparison_id for write in plan_a.writes)
    assert any(write.region_id == PROVENANCE_REGION and write.record_id == request.results[0].result_run_id
               for write in plan_a.writes)
    assert vba_sha256(first) == vba_sha256(second) == VBA_SHA


def test_t42_15_16_fixed_slot_gate38_gate39_candidate_compatibility(tmp_path):
    qualified = load_qualified_master(CANDIDATE_MANIFEST, CANDIDATE)
    run = core_result()
    comparison = run.comparisons[0]
    value = next(value for value in run.values if value.slice_id == comparison.left_slice_id)
    request = significance_qualification_request(qualified, run, comparison_id=comparison.comparison_id,
        value_id=value.value_id, token_id="canonical_marker", display_text="*")
    plan = plan_render(request, CANDIDATE)
    output = tmp_path / "fixed-candidate.xlsm"
    evidence = render(request, CANDIDATE, output, plan=plan)
    assert evidence["qa"] == "PASS" and evidence["vba_after"] == VBA_SHA
    workbook = load_workbook(output, keep_vba=True)
    try:
        assert workbook["presentation"]["E2"].value == "*"
        assert workbook["presentation"]["J2"].value is None
    finally:
        workbook.close(); workbook.vba_archive.close()


def test_t42_17_18_vba_formulas_names_and_protected_surfaces():
    assert vba_sha256(BASELINE) == vba_sha256(CANDIDATE) == VBA_SHA
    baseline = load_workbook(BASELINE, keep_vba=True, data_only=False)
    candidate = load_workbook(CANDIDATE, keep_vba=True, data_only=False)
    try:
        assert baseline["presentation"]["E1"].value == candidate["presentation"]["E1"].value == "=1+1"
        baseline_names = set(baseline.defined_names)
        assert baseline_names <= set(candidate.defined_names)
        assert set(candidate.defined_names) - baseline_names == {
            "ProductionDynamicResultAnchor", "ProductionDynamicProvenanceAnchor"}
        for sheet in baseline.sheetnames:
            if sheet != "presentation":
                assert tuple(tuple(cell.value for cell in row) for row in baseline[sheet].iter_rows()) == tuple(
                    tuple(cell.value for cell in row) for row in candidate[sheet].iter_rows())
    finally:
        baseline.close(); baseline.vba_archive.close()
        candidate.close(); candidate.vba_archive.close()
