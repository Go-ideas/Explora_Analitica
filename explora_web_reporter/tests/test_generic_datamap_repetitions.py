from __future__ import annotations

from io import BytesIO
from types import SimpleNamespace

import pandas as pd
import pytest
from openpyxl import load_workbook

from src.builder.analytic_db_builder import build_analytic_database
from src.builder.escalas_long_builder import build_escalas_long
from src.builder.respondentes_builder import build_respondentes
from src.database.db_reader import read_table
from src.export.datamap_exporter import save_final_datamap
from src.export.report_exporter import report_to_excel_bytes
from src.readers.datamap_reader import (
    get_review_source_sheet, normalize_review_datamap, read_datamap,
)
from src.reporter.tabulator import available_questions, generate_report


def metadata(labels):
    return SimpleNamespace(variable_value_labels=labels, column_names_to_labels={})


def test_suggested_types_fill_only_missing_canonical_type():
    raw = pd.DataFrame([
        {'variable': 'satisfaction', 'tipo_tabla_sugerida': 'Escala hedónica',
         'metrica_sugerida': 'Distribución, media; Top2Box 6-7 sugerido'},
        {'variable': 'choices', 'tipo_tabla_sugerida': 'RM % respondentes'},
        {'variable': 'fit', 'tipo_tabla_sugerida': 'JAR'},
        {'variable': 'count', 'tipo_tabla_sugerida': 'Escala 1-5'},
        {'variable': 'nominal', 'tipo_pregunta': 'RU', 'tipo_tabla_sugerida': 'Escala hedónica'},
        {'variable': 'unknown', 'tipo_tabla_sugerida': 'Modelo propietario sin definición'},
    ])
    review = normalize_review_datamap(raw).set_index('variable')
    assert review.loc['satisfaction', 'tipo_pregunta'] == 'Escala'
    assert review.loc['satisfaction', 'tipo_calculo'] == 'Media'
    assert review.loc['choices', 'tipo_pregunta'] == 'RM'
    assert review.loc['fit', 'tipo_pregunta'] == 'Escala JAR'
    assert review.loc['count', 'tipo_pregunta'] == 'Escala'
    assert review.loc['nominal', 'tipo_pregunta'] == 'RU'
    assert review.loc['unknown', 'tipo_pregunta'] == 'Sin clasificar'


def repeated_map(axis='Momento'):
    return pd.DataFrame([
        {'variable': 'rating_before', 'pregunta_id': 'SAT', 'texto_pregunta': 'Satisfacción',
         'tipo_tabla_sugerida': 'Escala hedónica', 'id_respondente': 'case_key',
         'Tipo_Eje_Fila': axis, 'Eje_Fila_Grid': 'Antes', 'Es_Item_Grid': 'Sí',
         'Grid_ID': 'SAT_REPETIDA', 'usar_en_dashboard': True},
        {'variable': 'rating_after', 'pregunta_id': 'SAT', 'texto_pregunta': 'Satisfacción',
         'tipo_tabla_sugerida': 'Escala hedónica', 'id_respondente': 'case_key',
         'Tipo_Eje_Fila': axis, 'Eje_Fila_Grid': 'Después', 'Es_Item_Grid': 'Sí',
         'Grid_ID': 'SAT_REPETIDA', 'usar_en_dashboard': True},
        {'variable': 'group', 'pregunta_id': 'GROUP', 'tipo_pregunta': 'RU',
         'clasificacion_analitica': 'Banner', 'es_banner': True, 'es_filtro': True,
         'id_respondente': 'case_key', 'usar_en_dashboard': True},
    ])


def test_repetitions_are_scalar_and_round_trip_is_stable(tmp_path):
    raw = repeated_map()
    review = normalize_review_datamap(raw)
    indexed = review.set_index('variable')
    assert indexed.at['rating_before', 'pregunta_id'] == 'SAT::rating_before'
    assert indexed.at['rating_after', 'pregunta_id'] == 'SAT::rating_after'
    assert 'Momento: Antes' in indexed.at['rating_before', 'texto_pregunta']
    assert indexed.at['rating_before', 'pregunta_padre'] == 'SAT'
    assert not indexed.at['rating_before', 'es_item_grid']
    path = tmp_path / 'final.xlsx'
    save_final_datamap(path, {'08_Clasificacion_Analitica': raw}, review)
    _, saved = get_review_source_sheet(read_datamap(path))
    restored = normalize_review_datamap(saved).set_index('variable')
    assert restored.pregunta_id.to_dict() == indexed.pregunta_id.to_dict()
    assert restored.texto_pregunta.to_dict() == indexed.texto_pregunta.to_dict()


@pytest.mark.parametrize('problem', ['missing_axis', 'missing_label', 'duplicate_label'])
def test_ambiguous_repetition_requires_validation(problem):
    raw = repeated_map()
    if problem == 'missing_axis': raw.loc[1, 'Tipo_Eje_Fila'] = ''
    elif problem == 'missing_label': raw.loc[1, 'Eje_Fila_Grid'] = ''
    else: raw.loc[1, 'Eje_Fila_Grid'] = 'Antes'
    with pytest.raises(ValueError, match='repetición'):
        normalize_review_datamap(raw)


def test_suffixes_alone_do_not_define_repetition():
    raw = pd.DataFrame({'variable': ['ATT.1', 'ATT.2'], 'pregunta_id': ['ATT', 'ATT'],
                        'tipo_pregunta': ['RM', 'RM']})
    review = normalize_review_datamap(raw)
    assert review.pregunta_id.tolist() == ['ATT', 'ATT']


def test_real_rm_options_keep_their_parent():
    raw = repeated_map()
    raw.loc[:1, 'tipo_tabla_sugerida'] = 'RM % respondentes'
    review = normalize_review_datamap(raw)
    assert review.pregunta_id.tolist()[:2] == ['SAT', 'SAT']


@pytest.mark.parametrize('low,high', [(1, 5), (1, 7), (1, 9), (0, 10), (-3, 3)])
def test_scale_uses_complete_labels_not_sample_extrema(low, high):
    df = pd.DataFrame({'score': [high-2, high-1, high-1]})
    datamap = normalize_review_datamap(pd.DataFrame([{'variable': 'score', 'pregunta_id': 'S', 'tipo_pregunta': 'Escala', 'usar_en_dashboard': True}]))
    people = build_respondentes(df, pd.DataFrame())
    scales = build_escalas_long(df, metadata({'score': {i: str(i) for i in range(low, high+1)}}), datamap, people)
    assert set(scales.escala_min) == {low}
    assert set(scales.escala_max) == {high}
    assert scales.top2box.tolist() == [0, 1, 1]
    # Ten-point scales are not automatically NPS.
    assert set(scales.nps_grupo) == {''}


def test_unknown_scale_does_not_fabricate_boxes_or_discard_zero(tmp_path):
    df = pd.DataFrame({'x': [0., 3., 4.]})
    review = normalize_review_datamap(pd.DataFrame([{'variable': 'x', 'tipo_pregunta': 'Escala', 'usar_en_dashboard': True}]))
    db = tmp_path / 'unknown.db'
    build_analytic_database(df, metadata({}), review, db_path=db, export_revision=False)
    report = generate_report(db, 'x', calculations=['Media', 'Top2Box'])
    assert report.calculations == ['Media']
    assert report.summary.media.iloc[0] == pytest.approx(7/3)
    assert any('límites de escala' in w for w in report.warnings)


def test_declared_scale_and_exclusions(tmp_path):
    df = pd.DataFrame({'score': [4., 5.], 'excluded': [4., 5.], 'rm': [1., 1.]})
    raw = pd.DataFrame([
        {'variable': 'score', 'tipo_pregunta': 'Escala', 'scale_min': 1, 'scale_max': 7, 'usar_en_dashboard': True},
        {'variable': 'excluded', 'tipo_pregunta': 'Escala', 'usar_en_dashboard': False},
        {'variable': 'rm', 'tipo_pregunta': 'RM', 'usar_en_dashboard': False},
    ])
    review = normalize_review_datamap(raw)
    db = tmp_path / 'explicit.db'
    build_analytic_database(df, metadata({}), review, db_path=db, export_revision=False)
    assert set(read_table(db, 'escalas_long').variable) == {'score'}
    assert read_table(db, 'multirrespuesta_long').empty
    r = generate_report(db, 'score', calculations=['Media', 'Top2Box'])
    assert r.summary.top2box.iloc[0] == 0


def test_declared_identity_is_separate_from_row_key():
    df = pd.DataFrame({'id': [10,11,12,13], 'case_key': ['a','a','b','b']})
    review = normalize_review_datamap(pd.DataFrame([{'variable': 'case_key', 'id_respondente': 'case_key'}]))
    people = build_respondentes(df, review)
    assert people.id_respondente.is_unique
    assert people.id_evaluacion.is_unique
    assert people.id_persona.tolist() == ['a','a','b','b']
    assert people.id_original.tolist() == ['a','a','b','b']


@pytest.mark.parametrize('declared', ['absent', 'case_key'])
def test_invalid_declared_identity_fails(declared):
    df = pd.DataFrame({'case_key': ['a', None]})
    dm = pd.DataFrame({'variable_id_persona': [declared]})
    with pytest.raises(ValueError, match='persona'):
        build_respondentes(df, dm)


def test_repeated_workflow_counts_people_and_blocks_independent_tests(tmp_path):
    df = pd.DataFrame({'case_key':['a','a','b','b'], 'group':[1,2,1,2],
                       'rating_before':[3,4,5,6], 'rating_after':[4,5,6,7]})
    review = normalize_review_datamap(repeated_map(), list(df))
    db = tmp_path/'repeated.db'
    labels = {var: {i:str(i) for i in range(1,8)} for var in ['rating_before','rating_after']}
    build_analytic_database(df, metadata(labels), review, db_path=db, export_revision=False)
    questions = available_questions(db)
    assert 'SAT_REPETIDA' not in set(questions.pregunta_id)
    assert 'SAT' not in set(questions.pregunta_id)
    assert {'SAT::rating_before','SAT::rating_after'}.issubset(set(questions.pregunta_id))
    report = generate_report(db, 'SAT::rating_before', banner=['group'], calculations=['Media','Top2Box'],
                             include_significance=True, min_base=2)
    assert report.question_type == 'Escala'
    assert report.base == 4 and report.person_base == 2
    assert report.base_unit == 'Evaluaciones'
    assert report.summary.loc[report.summary.banner.eq('Total'),'media'].iloc[0] == 4.5
    assert report.significance.empty and report.significance_display == 'none'
    assert any('muestras relacionadas' in warning for warning in report.warnings)
    assert any('4 evaluaciones de 2 personas' in note for note in report.notes)
    workbook = load_workbook(BytesIO(report_to_excel_bytes(report)), data_only=True)
    config = {row[0]:row[1] for row in workbook['Configuracion'].iter_rows(values_only=True) if len(row)>1}
    assert config['Unidad de base'] == 'Evaluaciones'
    assert str(config['Personas únicas']) == '2'
    filtered = generate_report(db, 'SAT::rating_after', filters={'group':[1]}, calculations=['Media'])
    assert filtered.base == 2 and filtered.person_base == 2
    assert filtered.summary.media.iloc[0] == 5


def test_scalar_variables_with_same_question_id_do_not_become_rm():
    raw = pd.DataFrame({'variable':['age','age_band'], 'pregunta_id':['AGE','AGE'],
                        'tipo_pregunta':['RU','RU']})
    review = normalize_review_datamap(raw)
    assert review.pregunta_id.tolist() == ['AGE::age','AGE::age_band']
    assert all('Variable:' in title for title in review.texto_pregunta)


@pytest.mark.parametrize('axis', ['Visita', 'Tienda', 'Sesión'])
def test_repetition_axes_are_not_specific_to_time(axis):
    review = normalize_review_datamap(repeated_map(axis))
    assert f'{axis}: Antes' in review.iloc[0].texto_pregunta


def test_repeated_menu_items_inherit_questionnaire_metadata(tmp_path):
    from src.utils.question_order import order_question_catalog
    from src.builder.preguntas_builder import build_preguntas
    review = normalize_review_datamap(repeated_map())
    questions = build_preguntas(review)
    path = tmp_path/'questionnaire.xlsx'
    with pd.ExcelWriter(path) as writer:
        pd.DataFrame([{'Pregunta_ID':'SAT','Orden_Cuestionario':1,
                       'Seccion_Cuestionario':'Experiencia', 'Texto_Pregunta':'Satisfacción'}]).to_excel(
            writer, sheet_name='13_Orden_Menu_Reporteador', index=False)
    ordered = order_question_catalog(questions, [path]).set_index('pregunta_id')
    assert not ordered.at['SAT::rating_before','es_variable_adicional']
    assert ordered.at['SAT::rating_before','seccion_reporter'] == 'Experiencia'
    assert 'Momento: Antes' in ordered.at['SAT::rating_before','menu_label']


def test_explicit_nps_uses_its_definition_and_keeps_zero():
    df = pd.DataFrame({'recommend':[0,7,10]})
    dm = normalize_review_datamap(pd.DataFrame([{'variable':'recommend','tipo_pregunta':'NPS','usar_en_dashboard':True}]))
    result = build_escalas_long(df, metadata({}), dm, build_respondentes(df,dm))
    assert result.nps_grupo.tolist() == ['Detractor','Pasivo','Promotor']
    assert result.valor.tolist() == [0,7,10]


@pytest.mark.parametrize('bounds', [(1,None), (5,1), (1,3)])
def test_invalid_or_inconsistent_scale_is_not_silently_used(bounds):
    df = pd.DataFrame({'score':[4,5]})
    dm = normalize_review_datamap(pd.DataFrame([{'variable':'score','tipo_pregunta':'Escala',
            'scale_min':bounds[0], 'scale_max':bounds[1], 'usar_en_dashboard':True}]))
    with pytest.raises(ValueError, match='escala'):
        build_escalas_long(df, metadata({}), dm, build_respondentes(df,dm))


def test_legacy_scale_fallback_uses_labels_and_is_not_automatically_nps():
    from src.reporter.tabulator import _scale_fallback
    data = pd.DataFrame({'valor_numerico':[0,7,8]})
    question = pd.Series({'pregunta_id':'Q','tipo_pregunta':'Escala'})
    options = pd.DataFrame({'pregunta_id':['Q']*11, 'codigo':range(11), 'label':[str(i) for i in range(11)]})
    result = _scale_fallback(data, question, options)
    assert result.valor.tolist() == [0,7,8]
    assert result.top2box.tolist() == [0,0,0]
    assert set(result.nps_grupo)=={''}
    unknown = _scale_fallback(data, question, pd.DataFrame())
    assert unknown.top2box.isna().all()


@pytest.mark.parametrize('unit', ['Respondente', 'Persona', 'respondent'])
def test_person_unit_counts_each_person_once(tmp_path, unit):
    df = pd.DataFrame({'person':['a','a','b','b','c','c','d','d'],
                       'gender':[1,1,1,1,2,2,2,2], 'site':[1,1,2,2,1,1,2,2]})
    raw = pd.DataFrame([
        {'variable':'gender','pregunta_id':'DEM','tipo_pregunta':'RU',
         'id_respondente':'person','Base_Valida_Nivel':unit,'usar_en_dashboard':True},
        {'variable':'site','pregunta_id':'SITE','tipo_pregunta':'RU',
         'id_respondente':'person','es_banner':True,'es_filtro':True,'usar_en_dashboard':True},
    ])
    dm = normalize_review_datamap(raw,list(df))
    db = tmp_path/'person.db'
    build_analytic_database(df,metadata({}),dm,db_path=db,export_revision=False)
    r = generate_report(db,'DEM',banner=['site'],calculations=['n','%'])
    assert r.base==4 and r.person_base==4 and r.base_unit=='Respondentes'
    assert sorted(r.summary.loc[r.summary.banner.eq('Total'),'n'])==[2,2]
    assert r.summary.loc[r.summary.banner.eq('Total'),'porcentaje'].sum()==1
    filtered=generate_report(db,'DEM',filters={'site':[1]},calculations=['n','%'])
    assert filtered.base==2 and filtered.base_before_filters==4


def test_conflicting_person_level_answers_require_validation(tmp_path):
    df = pd.DataFrame({'person':['a','a'], 'choice':[1,2]})
    dm = normalize_review_datamap(pd.DataFrame([{'variable':'choice','tipo_pregunta':'RU',
            'id_respondente':'person','unidad_analisis':'Persona','usar_en_dashboard':True}]),list(df))
    db=tmp_path/'conflict.db'
    build_analytic_database(df,metadata({}),dm,db_path=db,export_revision=False)
    with pytest.raises(ValueError, match='respuestas diferentes'):
        generate_report(db,'choice',calculations=['n','%'])


def test_changing_banner_is_not_arbitrarily_assigned_to_person(tmp_path):
    df=pd.DataFrame({'person':['a','a','b','b'],'answer':[1,1,2,2],'product':[1,2,1,2]})
    dm=normalize_review_datamap(pd.DataFrame([
        {'variable':'answer','tipo_pregunta':'RU','id_respondente':'person','unidad_analisis':'Persona','usar_en_dashboard':True},
        {'variable':'product','tipo_pregunta':'RU','id_respondente':'person','es_banner':True,'es_filtro':True,'usar_en_dashboard':True},
    ]),list(df))
    db=tmp_path/'banner.db'
    build_analytic_database(df,metadata({}),dm,db_path=db,export_revision=False)
    with pytest.raises(ValueError, match='cambia dentro de una persona'):
        generate_report(db,'answer',banner=['product'],calculations=['n'])
    r=generate_report(db,'answer',banner=['product'],filters={'product':[1]},calculations=['n'])
    assert r.base==2


def test_repeated_ids_export_to_portable_files_and_distinct_long_sheet_names(tmp_path):
    from zipfile import ZipFile
    from src.export.project_exporter import executive_export_bytes
    from src.export.report_exporter import reports_to_excel_bytes
    from dataclasses import replace
    df=pd.DataFrame({'x':[1,2]})
    dm=normalize_review_datamap(pd.DataFrame([{'variable':'x','tipo_pregunta':'RU','usar_en_dashboard':True}]))
    db=tmp_path/'export.db'
    build_analytic_database(df,metadata({}),dm,db_path=db,export_revision=False)
    report=generate_report(db,'x',calculations=['n','%'])
    report=replace(report,question_id='LONG_PARENT::very_long_physical_variable_name',
                   significance_display='separate',significance=pd.DataFrame({'Respuesta':['A'],'Grupo':['B']}))
    with ZipFile(BytesIO(executive_export_bytes(report))) as archive:
        assert all(':' not in name and '/' not in name for name in archive.namelist())
    workbook=load_workbook(BytesIO(reports_to_excel_bytes([report])),data_only=True)
    assert len(workbook.sheetnames)==2 and len(set(workbook.sheetnames))==2


def test_upload_structure_detection_cannot_merge_explicit_repeated_items():
    from src.structure.grid_loop_detector import enrich_datamap_with_spss_structure
    df = pd.DataFrame({'rating_before':[3,4], 'rating_after':[5,6], 'group':[1,2]})
    dm=normalize_review_datamap(repeated_map(),list(df))
    meta=metadata({c:{i:str(i) for i in range(1,8)} for c in ['rating_before','rating_after']})
    enriched=enrich_datamap_with_spss_structure(dm,df,meta)
    assert enriched.set_index('variable').pregunta_id.to_dict()==dm.set_index('variable').pregunta_id.to_dict()
    assert enriched.attrs['grid_loop_structures'].empty


def test_structure_detection_does_not_reenable_explicit_exclusions():
    from src.structure.grid_loop_detector import enrich_datamap_with_spss_structure
    df=pd.DataFrame({'GRID_1':[1,2], 'GRID_2':[2,3]})
    dm=normalize_review_datamap(pd.DataFrame({'variable':list(df),'pregunta_id':['GRID','GRID'],
                                           'tipo_pregunta':['Escala','Escala'],'usar_en_dashboard':[False,False]}))
    enriched=enrich_datamap_with_spss_structure(dm,df,metadata({}))
    assert not enriched.usar_en_dashboard.any()
    assert enriched.attrs['grid_loop_structures'].empty


def test_declared_rm_is_not_reclassified_from_question_wording():
    from src.structure.grid_loop_detector import enrich_datamap_with_spss_structure
    df=pd.DataFrame({'benefit_1':[1,0], 'benefit_2':[0,1]})
    dm=normalize_review_datamap(pd.DataFrame({'variable':list(df),'pregunta_id':['B','B'],
       'tipo_pregunta':['RM','RM'],'texto_pregunta':['¿Cuáles beneficios obtiene?']*2,'usar_en_dashboard':[True,True]}))
    enriched=enrich_datamap_with_spss_structure(dm,df,metadata({}))
    assert set(enriched.tipo_pregunta)=={'RM'}
    assert enriched.usar_en_dashboard.all()
    assert enriched.pregunta_id.tolist()==['B','B']
