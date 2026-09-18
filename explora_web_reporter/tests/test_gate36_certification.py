from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook

from scripts.gate36_certification import (
    check_interface, create_generic, formula_write, guarded_destination,
    inspect_package, literal_write,
    validate_vba_relationships,
    targeted_write,
)


@pytest.fixture
def generic(tmp_path):
    path = tmp_path / "generic.xlsx"
    create_generic(path)
    return path


def test_generic_interface(generic):
    check_interface(generic)


@pytest.mark.parametrize("text", ["=1+1", "+1", "-1", "@SUM(A1)", " =1", "\t=1"])
def test_literal_roundtrip(generic, text):
    w = load_workbook(generic)
    literal_write(w["presentation"]["F1"], text)
    w.save(generic)
    w.close()
    w = load_workbook(generic)
    try:
        assert w["presentation"]["F1"].value == text
        assert w["presentation"]["F1"].data_type == "s"
        assert w["presentation"]["E1"].data_type == "f"
    finally:
        w.close()


def test_allowlisted_formula(generic):
    w = load_workbook(generic)
    formula_write(w["presentation"]["E1"], "=1+1")
    w.save(generic)
    w.close()
    w = load_workbook(generic)
    assert w["presentation"]["E1"].value == "=1+1"
    assert w["presentation"]["E1"].data_type == "f"
    w.close()


def test_unauthorized_formula(generic):
    w = load_workbook(generic)
    with pytest.raises(ValueError, match="allowlisted"):
        formula_write(w["presentation"]["F1"], "=1+1")
    w.close()


@pytest.mark.parametrize("missing", ["role", "name", "table"])
def test_missing_interface_fails(generic, missing):
    w = load_workbook(generic)
    if missing == "role":
        del w["qa"]
    elif missing == "name":
        del w.defined_names["RendererSlot"]
    else:
        del w["presentation"].tables["CertificationResults"]
    w.save(generic)
    w.close()
    with pytest.raises(ValueError):
        check_interface(generic)


def test_xlsx_rejected_for_vba(generic):
    with pytest.raises(ValueError, match="requires .xlsm"):
        inspect_package(generic)


def test_missing_macro_type_rejected(generic):
    # This is a negative path, not a substitute genuine VBA fixture.
    path = generic.with_suffix(".xlsm")
    path.write_bytes(generic.read_bytes())
    with pytest.raises(ValueError, match="content type"):
        inspect_package(path)


def test_corrupt_package(tmp_path):
    path = tmp_path / "corrupt.xlsm"
    path.write_bytes(b"not a zip")
    with pytest.raises(ValueError, match="Invalid workbook"):
        inspect_package(path)


def test_missing_vba_payload(generic):
    path = generic.with_name("negative_missing_payload.xlsm")
    with ZipFile(generic) as source, ZipFile(path, "w") as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(
                    b"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml",
                    b"application/vnd.ms-excel.sheet.macroEnabled.main+xml")
            target.writestr(item, data)
    with pytest.raises(ValueError, match="Invalid workbook"):
        inspect_package(path)


def test_missing_vba_relationship():
    with pytest.raises(ValueError, match="relationship"):
        validate_vba_relationships(ET.Element("Relationships"))


@pytest.mark.parametrize("target,external", [("wrong.bin", False), ("vbaProject.bin", True)])
def test_invalid_vba_relationship(target, external):
    from scripts.gate36_certification import VBA_REL
    rels = ET.Element("Relationships")
    attrs = {"Type": VBA_REL, "Target": target}
    if external:
        attrs["TargetMode"] = "External"
    ET.SubElement(rels, "Relationship", attrs)
    with pytest.raises(ValueError, match="relationship"):
        validate_vba_relationships(rels)


def test_generic_authorized_value_update(generic):
    before = inspect_package(generic, require_vba=False)
    w = load_workbook(generic)
    w["presentation"]["B2"] = 0.25
    w.save(generic)
    w.close()
    after = inspect_package(generic, require_vba=False)
    for name in before["groups"]["tables"] + ["xl/styles.xml"]:
        assert before["parts"][name]["semantic"] == after["parts"][name]["semantic"]
    assert before["formulas"] == after["formulas"]
    assert before["defined_names"] == after["defined_names"]


@pytest.mark.parametrize("text", ["=1+1", "+1", "-1", "@SUM(A1)", " =1", "\t=1"])
def test_targeted_literal_path(generic, text):
    destination = generic.with_name("targeted.xlsx")
    targeted_write(generic, destination, "F1", text, require_vba=False)
    w = load_workbook(destination)
    assert w["presentation"]["F1"].value == text
    assert w["presentation"]["F1"].data_type == "s"
    w.close()


def test_targeted_formula_path(generic):
    destination = generic.with_name("targeted.xlsx")
    targeted_write(generic, destination, "E1", "=1+1", kind="formula", require_vba=False)
    w = load_workbook(destination)
    assert w["presentation"]["E1"].value == "=1+1"
    assert w["presentation"]["E1"].data_type == "f"
    w.close()


def test_targeted_numeric_and_preserved_parts(generic):
    destination = generic.with_name("targeted.xlsx")
    targeted_write(generic, destination, "B2", 0.25, kind="numeric", require_vba=False)
    before = inspect_package(generic, require_vba=False)
    after = inspect_package(destination, require_vba=False)
    assert all(before["parts"][n]["sha256"] == after["parts"][n]["sha256"]
               for n in before["parts"] if n != "xl/worksheets/sheet1.xml")
    w = load_workbook(destination)
    assert w["presentation"]["B2"].value == 0.25
    assert w["presentation"]["B2"].number_format == "0.0%"
    w.close()


@pytest.mark.parametrize("address,kind,value", [("D1", "literal", "overwrite"),
    ("E1", "formula", "=SUM(A1:A2)"), ("B2", "numeric", float("nan"))])
def test_targeted_rejects_unauthorized_operations(generic, address, kind, value):
    destination = generic.with_name("targeted.xlsx")
    with pytest.raises(ValueError):
        targeted_write(generic, destination, address, value, kind=kind, require_vba=False)
    assert not destination.exists()


def test_targeted_missing_interface(generic):
    w = load_workbook(generic)
    del w["qa"]
    w.save(generic)
    w.close()
    destination = generic.with_name("targeted.xlsx")
    with pytest.raises(ValueError):
        targeted_write(generic, destination, "F1", "literal", require_vba=False)
    assert not destination.exists()


def test_targeted_determinism(generic):
    a, b = generic.with_name("a.xlsx"), generic.with_name("b.xlsx")
    for destination in (a, b):
        targeted_write(generic, destination, "B2", 0.25, kind="numeric", require_vba=False)
    left, right = inspect_package(a, require_vba=False), inspect_package(b, require_vba=False)
    assert {n:p["sha256"] for n,p in left["parts"].items()} == {n:p["sha256"] for n,p in right["parts"].items()}


@pytest.mark.parametrize("role", ["provenance", "results", "bases", "slices", "significance", "qa", "cell_map", "canonical_payload"])
def test_targeted_registered_audit_role(generic, role):
    destination = generic.with_name("role_output.xlsx")
    targeted_write(generic, destination, "B2", "GENERIC_TRANSPORT_ONLY",
                   kind="literal", require_vba=False, role=role)
    check_interface(destination)
    w = load_workbook(destination)
    assert w[role]["B2"].value == "GENERIC_TRANSPORT_ONLY"
    assert w[role]["B2"].data_type == "s"
    w.close()


@pytest.mark.parametrize("role", ["configuration", "navigation"])
def test_targeted_master_role_rejected(generic, role):
    with pytest.raises(ValueError, match="Master-owned"):
        targeted_write(generic, generic.with_name("role_output.xlsx"), "B2", "overwrite",
                       kind="literal", require_vba=False, role=role)


@pytest.mark.parametrize("mode,suffix", [(False, ".xlsm"), (True, ".xlsx")])
def test_invalid_destination_or_mode(generic, mode, suffix):
    with pytest.raises(ValueError, match="Preservation mode"):
        guarded_destination(generic, generic.with_name("output" + suffix), preserve_vba=mode)


def test_generic_determinism_and_ownership(generic):
    before = inspect_package(generic, require_vba=False)
    w = load_workbook(generic)
    s = w["presentation"]
    assert s["D1"].value == "MASTER_OWNED"
    assert s["D2"].value == "VBA_OWNED_CONCEPT_ONLY_NO_MACRO"
    assert s["D3"].protection.locked is False
    assert s["B2"].number_format == "0.0%"
    assert w["results"].sheet_state == "hidden"
    w.save(generic)
    w.close()
    after = inspect_package(generic, require_vba=False)
    assert before["sheets"] == after["sheets"]
    assert before["defined_names"] == after["defined_names"]
    assert before["formulas"] == after["formulas"]
    assert before["groups"]["tables"] == after["groups"]["tables"]
    assert {n: p["semantic"] for n, p in before["parts"].items()} == {
        n: p["semantic"] for n, p in after["parts"].items()}
