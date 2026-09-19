from __future__ import annotations

from dataclasses import dataclass, fields
from decimal import Decimal, ROUND_HALF_EVEN
import hashlib
import json
import math
from pathlib import Path
import posixpath
import shutil
import sys
from tempfile import TemporaryDirectory
from zipfile import ZipFile

from lxml import etree
from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries

from scripts.gate36_certification import inspect_package
from src.analytics_core.result_identity import (
    canonical_payload, fingerprint, request_identity, result_fingerprint, stable_json,
)
from src.analytics_core.serialization import from_canonical_json, to_canonical_data, to_canonical_json
from src.contracts.models import (
    CanonicalBase, CanonicalResult, CanonicalSlice, CanonicalValue, ReleaseMetadata,
    SignificanceRelation,
)
from src.contracts.validators import validate_canonical_result, validate_release_metadata
from src.contracts.vocabulary import StatisticalState, ValueStatus, ValueUnit

VISUAL = "M7_VISUAL_SPEC_V1"
MASTER = "M7_MASTER_INTERFACE_CONTRACT_V1"
MAPPING = "M7_CANONICAL_TO_WORKBOOK_MAPPING_V1"
NUMERIC = "M7_EXCEL_NUMERIC_DISPLAY_PROFILE_V1"
SIGNIFICANCE = "M7_SIGNIFICANCE_PRESENTATION_INTERFACE_V1"
BUILD = "M7B_RENDERER_V1"
S = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
ROLES = ("presentation", "provenance", "results", "bases", "slices", "significance",
         "qa", "cell_map", "canonical_payload", "configuration", "navigation")
COLLECTIONS = {"value": "values", "base": "bases", "slice": "slices",
               "comparison": "comparisons", "qa": "qa_events"}
ID_FIELDS = {"value": "value_id", "base": "base_id", "slice": "slice_id",
             "comparison": "comparison_id", "qa": "qa_id"}


class RenderError(ValueError):
    pass


def require(condition, message):
    if not condition:
        raise RenderError(message)


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def text(value, name):
    require(isinstance(value, str) and bool(value.strip()), f"Missing identity: {name}")


def literal(value):
    require(isinstance(value, str), "Literal text required")
    require(len(value.encode("utf-16-le")) // 2 <= 32767, "Text capacity exceeded")
    try:
        etree.Element("text").text = value
    except (ValueError, UnicodeError) as exc:
        raise RenderError("Invalid XML text") from exc
    return value


@dataclass(frozen=True)
class TableBinding:
    role_id: str
    sheet: str
    table: str
    capacity: int


@dataclass(frozen=True)
class Slot:
    slot_id: str
    role_id: str
    kind: str
    storage_mode: str
    # Exactly one name or bounded existing table-body field, never coordinates.
    name: str | None = None
    table: str | None = None
    row: int | None = None
    column: str | None = None
    status_name: str | None = None
    significance_token_id: str | None = None
    significance_value_id: str | None = None
    label_name: str | None = None


@dataclass(frozen=True)
class MasterManifest:
    master_id: str
    version: str
    artifact_sha256: str
    interface_version: str
    compatible_visual_spec_versions: tuple[str, ...]
    roles: tuple[TableBinding, ...]
    writable_slots: tuple[Slot, ...] = ()
    # Every existing cell formula must have an exact presentation-only classification.
    presentation_formulas: tuple[tuple[str, str, str], ...] = ()


@dataclass(frozen=True)
class RenderRequest:
    results: tuple[CanonicalResult, ...]
    visual_spec_json: str
    master: MasterManifest
    slots: tuple[Slot, ...]
    configuration_id: str
    configuration_version: str
    configuration_release: ReleaseMetadata
    master_release: ReleaseMetadata
    mapping_version: str = MAPPING
    significance_json: tuple[str, ...] = ()
    renderer_build: str = BUILD


@dataclass(frozen=True)
class PlannedWrite:
    sheet: str
    cell: str
    kind: str
    value: str | int | float | None
    number_format: str
    result_run_id: str
    result_fingerprint: str
    record_role: str
    record_id: str
    field_name: str
    slot_id: str
    visual_id: str
    display_profile: str
    storage_mode: str
    status: str
    unit: str
    token_id: str | None = None


@dataclass(frozen=True)
class RenderPlan:
    writes: tuple[PlannedWrite, ...]
    input_sha256: str
    vba_sha256: str
    canonical_json: tuple[str, ...]
    provenance_json: str
    request_sha256: str
    plan_sha256: str


def configuration_fingerprint(request):
    return fingerprint({"visual_spec": request.visual_spec_json, "master": request.master,
                        "slots": request.slots, "tokens": request.significance_json,
                        "mapping": request.mapping_version,
                        "configuration_id": request.configuration_id,
                        "configuration_version": request.configuration_version,
                        "build": request.renderer_build})


def decode(raw):
    def pairs(items):
        out = {}
        for key, value in items:
            require(key not in out, "Duplicate JSON field")
            out[key] = value
        return out
    try:
        return json.loads(raw, object_pairs_hook=pairs,
                          parse_constant=lambda v: require(False, "Nonfinite JSON"))
    except (ValueError, TypeError) as exc:
        raise RenderError("Malformed contract JSON") from exc


def shape(obj, required, optional=()):
    require(isinstance(obj, dict), "Contract object required")
    require(set(required) <= set(obj) <= set(required) | set(optional),
            "Unknown or missing contract fields")


def array(value, name, nonempty=False):
    require(isinstance(value, list) and (not nonempty or bool(value)), f"Invalid array: {name}")


def label(obj):
    shape(obj, ("text", "language"), ("source_ref",))
    literal(obj["text"])
    text(obj["language"], "language")
    if "source_ref" in obj:
        text(obj["source_ref"], "label source")


def visual_spec(request, runs):
    spec = decode(request.visual_spec_json)
    shape(spec, ("schema_version", "visual_spec_id", "revision", "project_id", "result_refs",
                 "master_interface_version", "numeric_profile_version", "sections"), ("provenance_refs",))
    spec.setdefault("provenance_refs", [])
    require(spec["schema_version"] == VISUAL and spec["master_interface_version"] == MASTER
            and spec["numeric_profile_version"] == NUMERIC, "Unsupported contract version")
    text(spec["visual_spec_id"], "visual_spec_id")
    require(type(spec["revision"]) is int and spec["revision"] > 0, "Invalid revision")
    require({r.project_id for r in runs.values()} == {spec["project_id"]}, "Cross-project reference")
    array(spec["result_refs"], "result_refs", True)
    array(spec["provenance_refs"], "provenance_refs")
    for ref in spec["provenance_refs"]:
        text(ref, "provenance reference")
    seen = set()
    for ref in spec["result_refs"]:
        shape(ref, ("result_run_id", "result_fingerprint", "request_fingerprint"))
        run = runs.get(ref["result_run_id"])
        require(run is not None and run.result_run_id not in seen, "Unknown or duplicate run reference")
        require(ref["result_fingerprint"] == run.result_fingerprint
                and ref["request_fingerprint"] == run.request.request_fingerprint, "Source identity mismatch")
        seen.add(run.result_run_id)
    require(seen == set(runs), "Incomplete result references")
    ids = set()
    array(spec["sections"], "sections", True)
    for section in spec["sections"]:
        shape(section, ("section_id", "order", "sheet_role", "title", "visuals"))
        label(section["title"])
        require(section["sheet_role"] in ROLES[:9], "Unknown sheet role")
        array(section["visuals"], "visuals", True)
        for obj, id_field in [(section, "section_id"), *[(v, "visual_id") for v in section["visuals"]]]:
            text(obj.get(id_field), id_field)
            require(obj[id_field] not in ids, "Duplicate visual/section ID")
            ids.add(obj[id_field])
            require(type(obj.get("order")) is int and obj["order"] >= 0, "Invalid order")
        for visual in section["visuals"]:
            shape(visual, ("visual_id", "order", "role", "source_bindings", "labels",
                           "display_profile_ref"), ("warning_refs", "significance_presentation_refs"))
            visual.setdefault("warning_refs", [])
            visual.setdefault("significance_presentation_refs", [])
            require(visual["role"] in ("result_table", "base_table", "significance_table", "qa_table", "provenance"), "Unsupported visual")
            array(visual["source_bindings"], "source_bindings", True)
            array(visual["labels"], "labels")
            for l in visual["labels"]:
                label(l)
            for key in ("warning_refs", "significance_presentation_refs"):
                array(visual[key], key)
            known_qa = {q.qa_id for r in runs.values() for q in r.qa_events}
            require(set(visual["warning_refs"]) <= known_qa, "Unknown warning reference")
    return spec


def validate_sources(request):
    require(bool(request.results), "Missing canonical result")
    runs = {}
    for run in request.results:
        require(isinstance(run, CanonicalResult), "CanonicalResult required")
        # Own a canonical snapshot: upstream dataclasses can contain mutable nested dictionaries.
        run = from_canonical_json(to_canonical_json(run))
        validate_canonical_result(run)
        require(run.result_schema_version == "M5_CANONICAL_RESULT_V1", "Unsupported canonical schema")
        require(run.manifest is not None and run.request is not None and run.release is not None,
                "Incomplete canonical provenance")
        require(run.release.releasable and str(run.release.computation_status) == "COMPLETE",
                "Canonical release not eligible")
        require(str(run.release.qa_release_status) in ("PASS", "PASS_WITH_WARNINGS"),
                "Contradictory canonical release state")
        require(not any(q.blocking for q in run.qa_events)
                and str(run.qa.aggregate_state) != "FAIL", "Blocking QA")
        require(run.result_fingerprint == result_fingerprint(run)
                == run.manifest.result_fingerprint, "Canonical fingerprint mismatch")
        require(run.manifest.schema_version == run.result_schema_version
                and run.manifest.core_version == run.core_version, "Manifest identity mismatch")
        text(run.manifest.ruleset, "ruleset")
        require(run.project_spec_ref in run.manifest.source_spec_refs, "Missing source spec reference")
        _, request_fp = request_identity(canonical_payload(run.request))
        require(run.request.request_fingerprint == request_fp, "Request fingerprint mismatch")
        require(run.result_run_id not in runs, "Ambiguous canonical run")
        slices = {s.slice_id: s for s in run.slices}
        bases = {b.base_id: b for b in run.bases}
        for b in run.bases:
            require(b.slice_id in slices, "Unresolved base slice")
        for v in run.values:
            ValueStatus(str(v.value_status))
            ValueUnit(str(v.unit))
            b = bases[v.base_id]
            require(v.slice_id in slices and (v.question_id, v.structure_id, v.slice_id)
                    == (b.question_id, b.structure_id, b.slice_id), "Contradictory value/base identity")
        for c in run.comparisons:
            StatisticalState(str(c.status))
            require(c.left_slice_id in slices and c.right_slice_id in slices,
                    "Unresolved comparison slice")
        runs[run.result_run_id] = run
    require(len({(r.project_id, r.dataset_fingerprint, r.project_spec_ref) for r in runs.values()}) == 1,
            "Mixed canonical projects/datasets")
    return runs


def open_master(path, request):
    manifest = request.master
    for key in ("master_id", "version", "artifact_sha256"):
        text(getattr(manifest, key), key)
    require(manifest.interface_version == MASTER and VISUAL in manifest.compatible_visual_spec_versions,
            "Incompatible Master version")
    require(sha(path) == manifest.artifact_sha256, "Master hash mismatch")
    require(bool(manifest.writable_slots) and request.slots == manifest.writable_slots,
            "Undeclared target / incomplete Master slot coverage")
    inventory = inspect_package(Path(path))
    require(inventory["whole_sha256"] == manifest.artifact_sha256, "Master changed during inspection")
    with ZipFile(path) as z:
        require(not any("signature" in n.lower() for n in z.namelist()), "Signed package unsupported")
        require(all(p.file_size <= 64 * 1024 * 1024 for p in z.infolist()
                    if p.filename.endswith(".xml")), "Master XML resource bound exceeded")
    w = load_workbook(path, keep_vba=True, keep_links=True)
    try:
        roles = {r.role_id: r for r in manifest.roles}
        require(w.calculation is None or w.calculation.fullPrecision is not False,
                "Precision-as-displayed would alter canonical scalars")
        require(len(roles) == len(manifest.roles) and set(roles) == set(ROLES), "Missing/ambiguous role registry")
        require(len({r.sheet for r in roles.values()}) == len(ROLES), "Ambiguous role sheets")
        require(set(w.sheetnames) == {r.sheet for r in roles.values()}, "Unknown sheet ownership")
        for role, binding in roles.items():
            require(binding.sheet in w.sheetnames, "Missing Master sheet")
            sheet = w[binding.sheet]
            require(type(binding.capacity) is int and binding.capacity > 0, "Invalid capacity")
            require(binding.table in sheet.tables, "Missing declared table")
            bounds = range_boundaries(sheet.tables[binding.table].ref)
            require(bounds[3] - bounds[1] == binding.capacity, "Table capacity mismatch")
            table = sheet.tables[binding.table]
            require(table.headerRowCount == 1 and not table.totalsRowCount, "Unsupported table layout")
            require(not any(c.calculatedColumnFormula or c.totalsRowFormula for c in table.tableColumns),
                    "Unclassified table formulas")
            headers = [sheet.cell(bounds[1], i).value for i in range(bounds[0], bounds[2]+1)]
            require(headers == [c.name for c in table.tableColumns], "Table header binding mismatch")
            if role in ("results", "bases", "slices", "significance", "qa", "cell_map", "canonical_payload"):
                require(sheet.protection.sheet, "Audit role not protected")
            if role in ("presentation", "provenance", "configuration", "navigation"):
                require(sheet.sheet_state == "visible", "Required visible role hidden")
        config = roles["configuration"]
        registry = w[config.sheet].tables[config.table]
        columns = [c.name for c in registry.tableColumns]
        required_columns = ("role_id", "sheet", "owner", "slot_id", "table_binding", "capacity", "policy")
        require(len(columns) == len(required_columns) and set(columns) == set(required_columns),
                "Invalid role registry schema")
        left, top, right, bottom = range_boundaries(registry.ref)
        rows = [dict(zip(columns, row)) for row in w[config.sheet].iter_rows(
            min_row=top+1, max_row=bottom, min_col=left, max_col=right, values_only=True)]
        require(len(rows) == len(ROLES) and {r["role_id"] for r in rows} == set(ROLES), "Invalid physical role registry")
        for row in rows:
            role, sheet, owner, slot, table, capacity, policy = (row[k] for k in required_columns)
            expected = roles[role]
            require((sheet, table, capacity) == (expected.sheet, expected.table, expected.capacity),
                    "Physical registry mismatch")
            require(owner == ("MASTER" if role in ("configuration", "navigation") else "RENDERER"),
                    "Role ownership mismatch")
            named_cell(w, slot, sheet)
            require(policy == ("preserve" if owner == "MASTER" else "replace_declared_only"),
                    "Unknown update policy")
        formulas = {(s.title, c.coordinate, c.value) for s in w for row in s for c in row if c.data_type == "f"}
        require(formulas == set(manifest.presentation_formulas), "Unclassified formula inventory")
        require(all(n.type in ("RANGE", "STRING", "NUMBER", "BOOL") for n in w.defined_names.values()),
                "Unclassified named formula")
        return w, roles, inventory
    except Exception:
        w.close()
        w.vba_archive.close()
        raise


def named_cell(w, name, expected_sheet):
    require(name in w.defined_names, "Missing workbook target")
    destinations = list(w.defined_names[name].destinations)
    require(len(destinations) == 1, "Ambiguous target")
    sheet, address = destinations[0]
    require(sheet == expected_sheet, "Target role mismatch")
    address = address.replace("$", "")
    left, top, right, bottom = range_boundaries(address)
    require(left == right and top == bottom, "Ambiguous multi-cell target")
    return w[sheet].cell(top, left)


def resolve_slot(w, roles, slot):
    require(slot.role_id == "presentation", "Undeclared/owned target")
    binding = roles[slot.role_id]
    require(bool(slot.name) != bool(slot.table), "Ambiguous target binding")
    if slot.name:
        require(slot.row is None and slot.column is None, "Ambiguous target")
        cell = named_cell(w, slot.name, binding.sheet)
    else:
        require(slot.table == binding.table and type(slot.row) is int
                and 0 <= slot.row < binding.capacity, "Undeclared table target")
        table = w[binding.sheet].tables[slot.table]
        columns = [c.name for c in table.tableColumns]
        require(columns.count(slot.column) == 1, "Missing/ambiguous target column")
        left, top, _, _ = range_boundaries(table.ref)
        cell = w[binding.sheet].cell(top + 1 + slot.row, left + columns.index(slot.column))
    require(cell.data_type != "f", "Protected formula overwrite")
    require(cell.protection.locked, "User-owned target overlap")
    require(slot.kind in ("numeric", "literal") and slot.storage_mode in ("exact_numeric", "exact_text"),
            "Target type mismatch")
    require((slot.kind, slot.storage_mode) in (("numeric", "exact_numeric"), ("literal", "exact_text")),
            "Incompatible target storage mode")
    return cell


def present(value, record, field, profile, mode):
    if profile == "LITERAL":
        require(isinstance(value, str) and mode == "exact_text", "Target type mismatch")
        return "literal", literal(value), "General", "TEXT", "TEXT"
    require(field in ("estimate", "unweighted_n", "weighted_n", "weighted_n_raw"),
            "Unsupported numeric presentation binding")
    status = str(record.get("value_status", record.get("base_status", "OK")))
    allowed_statuses = {s.value for s in ValueStatus}
    if "base_status" in record:
        allowed_statuses |= {"VALID_ZERO_BASE", "NONZERO_BASE"}
    require(status in allowed_statuses, "Unknown source status")
    unit = str(record.get("unit", "COUNT"))
    if field == "effective_n":
        raise RenderError("effective_n is QA-only")
    count_integral = type(value) is int or (type(value) is float and math.isfinite(value)
                        and Decimal(stable_json(value)) == Decimal(stable_json(value)).to_integral_value())
    formats = {"PROPORTION": "0.0%", "MEAN": "0.00", "STANDARD_DEVIATION": "0.00",
               "SCORE": "0.00", "weighted_n": "0.00", "weighted_n_raw": "0.00",
               "unweighted_n": "0", "COUNT": "0" if count_integral else "0.00"}
    expected_profile = field if field in ("weighted_n", "weighted_n_raw", "unweighted_n") else unit
    require(profile in formats and profile == expected_profile, "Unknown/incompatible display profile")
    if status not in ("OK", "VALID_ZERO_BASE", "NONZERO_BASE") or value is None:
        fmt = "General" if mode == "exact_text" else formats[profile]
        return "blank", None, fmt, "NULL" if value is None and status in ("OK", "VALID_ZERO_BASE", "NONZERO_BASE") else status, unit
    require(type(value) in (int, float), "Numeric scalar required")
    require(type(value) is int or math.isfinite(value), "Nonfinite numeric scalar")
    if mode == "exact_text":
        return "literal", stable_json(value), "General", status, unit
    require(mode == "exact_numeric", "Unknown storage mode")
    if type(value) is float:
        require(value == 0 or abs(value) >= sys.float_info.min, "Excel denormal scalar unsupported")
        require(value != 0 or math.copysign(1, value) > 0, "Excel negative zero unsupported")
    decimal = Decimal(stable_json(value))
    require(len(decimal.normalize().as_tuple().digits) <= 15
            and (type(value) is not int or abs(value) <= 999999999999999), "Unsafe numeric precision")
    return "numeric", value, formats[profile], status, unit


def invariant_display(value, profile):
    """Frozen half-even text profile; never used as analytical authority."""
    require(profile in ("PROPORTION", "COUNT", "MEAN", "STANDARD_DEVIATION", "SCORE",
                        "unweighted_n", "weighted_n", "weighted_n_raw"), "Unknown display profile")
    decimal = Decimal(stable_json(value))
    if profile == "PROPORTION":
        return str((decimal * 100).quantize(Decimal("0.1"), rounding=ROUND_HALF_EVEN)) + "%"
    places = Decimal("1") if profile == "unweighted_n" or (profile == "COUNT" and decimal == decimal.to_integral_value()) else Decimal("0.01")
    return str(decimal.quantize(places, rounding=ROUND_HALF_EVEN))


def tokens(request, runs, spec):
    result = {}
    checksums = {sha_bytes(raw): raw for raw in request.significance_json}
    require(len(checksums) == len(request.significance_json), "Duplicate token envelope")
    refs = {ref for s in spec["sections"] for v in s["visuals"] for ref in v["significance_presentation_refs"]}
    require(refs == set(checksums), "Unknown/unpinned significance envelope")
    for raw in request.significance_json:
        envelope = decode(raw)
        shape(envelope, ("schema_version", "token_set_id", "revision", "authority_ref", "result_run_id",
                         "result_fingerprint", "comparison_bindings", "legend"))
        require(envelope["schema_version"] == SIGNIFICANCE, "Unsupported token contract")
        text(envelope["authority_ref"], "token authority")
        text(envelope["token_set_id"], "token_set_id")
        require(type(envelope["revision"]) is int and envelope["revision"] > 0, "Invalid token revision")
        run = runs.get(envelope["result_run_id"])
        require(run is not None and run.result_fingerprint == envelope["result_fingerprint"], "Token run mismatch")
        array(envelope["comparison_bindings"], "comparison_bindings", True)
        array(envelope["legend"], "legend", True)
        legend = {}
        for item in envelope["legend"]:
            shape(item, ("token_id", "meaning", "comparison_refs"))
            text(item["token_id"], "token_id")
            literal(item["meaning"])
            array(item["comparison_refs"], "comparison_refs", True)
            require(item["token_id"] not in legend, "Ambiguous token legend")
            legend[item["token_id"]] = item
        comparisons = {c.comparison_id: canonical_payload(c) for c in run.comparisons}
        require(all(set(item["comparison_refs"]) <= set(comparisons) for item in legend.values()),
                "Unresolved legend comparison")
        values = {v.value_id: canonical_payload(v) for v in run.values}
        for binding in envelope["comparison_bindings"]:
            shape(binding, ("comparison_id", "family_id", "test_id", "test_version", "value_id",
                            "slice_id", "member_id", "token_id", "display_text"), ("direction",))
            c = comparisons.get(binding["comparison_id"])
            v = values.get(binding["value_id"])
            require(c is not None and v is not None and c["status"] == "SIGNIFICANT", "Invalid significance assertion")
            require(all(binding[k] == c[k] for k in ("family_id", "test_id", "test_version"))
                    and binding.get("direction") == c["direction"], "Contradictory token identity/direction")
            require((binding["slice_id"], binding["member_id"]) in
                    ((c["left_slice_id"], c["left_member_id"]), (c["right_slice_id"], c["right_member_id"])), "Token target mismatch")
            require(v["slice_id"] == binding["slice_id"] and v["question_id"] == c["question_id"]
                    and v["metric_id"] == c["metric_id"] and v["value_status"] == "OK", "Token value mismatch")
            require(binding["token_id"] in legend and c["comparison_id"] in legend[binding["token_id"]]["comparison_refs"], "Unknown significance token")
            key = (run.result_run_id, sha_bytes(raw), binding["comparison_id"],
                   binding["value_id"], binding["token_id"])
            require(key not in result, "Ambiguous significance token")
            literal(binding["display_text"])
            result[key] = {**binding, "envelope_checksum": sha_bytes(raw)}
    return result


def plan_render(request: RenderRequest, master_path: Path) -> RenderPlan:
    require(request.mapping_version == MAPPING and request.renderer_build == BUILD, "Unsupported mapping/build")
    text(request.configuration_id, "configuration_id")
    text(request.configuration_version, "configuration_version")
    validate_release_metadata(request.configuration_release)
    validate_release_metadata(request.master_release)
    config_sha = configuration_fingerprint(request)
    require(request.configuration_release.source_hash == config_sha
            and request.master_release.source_hash == request.master.artifact_sha256, "Release checksum mismatch")
    runs = validate_sources(request)
    spec = visual_spec(request, runs)
    token_map = tokens(request, runs, spec)
    require(bool(request.slots), "Missing mapping")
    slots = {s.slot_id: s for s in request.slots}
    require(len(slots) == len(request.slots), "Ambiguous mapping")
    for slot in request.slots:
        text(slot.slot_id, "slot_id")
    w, roles, inventory = open_master(master_path, request)
    try:
        with ZipFile(master_path) as package:
            physical_cells = {}
            physical_children = {}
            parser = etree.XMLParser(resolve_entities=False, no_network=True)
            for sheet, part in worksheet_parts(package).items():
                tree = etree.fromstring(package.read(part), parser)
                addresses = [c.get("r") for c in tree.findall(f"{{{S}}}sheetData/{{{S}}}row/{{{S}}}c")]
                require(len(addresses) == len(set(addresses)), "Ambiguous physical cells")
                physical_cells[sheet] = set(addresses)
                physical_children[sheet] = {c.get("r"): {child.tag for child in c}
                    for c in tree.findall(f"{{{S}}}sheetData/{{{S}}}row/{{{S}}}c")}
    except Exception:
        w.close()
        w.vba_archive.close()
        raise
    writes = {}
    presentation = []
    renderer_warnings = []
    used_slots = set()

    def put(cell, kind, value, run, role, record_id, field, slot_id, visual_id="", profile="AUDIT_JSON",
            mode="exact_text", status="OK", unit="TEXT", token_id=None, number_format="General"):
        require(cell.data_type != "f", "Protected formula overwrite")
        require(cell.protection.locked, "User-owned write overlap")
        require(cell.coordinate in physical_cells.get(cell.parent.title, ()), "Missing physical target")
        require(physical_children[cell.parent.title][cell.coordinate] <= {f"{{{S}}}v", f"{{{S}}}is"},
                "Unsupported target child structure")
        require(kind != "literal" or literal(value) == value, "Invalid literal")
        key = (cell.parent.title, cell.coordinate)
        require(key not in writes, "Duplicate/ambiguous target")
        require(number_format == cell.number_format, "Target display format mismatch")
        writes[key] = PlannedWrite(*key, kind, value, number_format,
                                  run.result_run_id if run else "ALL_RUNS",
                                  run.result_fingerprint if run else fingerprint([r.result_fingerprint for r in runs.values()]),
                                  role, record_id, field, slot_id, visual_id, profile, mode, status, unit, token_id)

    try:
        for section in sorted(spec["sections"], key=lambda s: (s["order"], s["section_id"])):
            for visual in sorted(section["visuals"], key=lambda v: (v["order"], v["visual_id"])):
                for binding_index, binding in enumerate(visual["source_bindings"]):
                    shape(binding, ("result_run_id", "record_role", "record_id", "field_name", "slot_id"))
                    run = runs.get(binding["result_run_id"])
                    require(run is not None, "Missing canonical result")
                    slot = slots.get(binding["slot_id"])
                    require(slot is not None, "Missing/unknown mapping")
                    require(slot.role_id == section["sheet_role"], "Mapping role mismatch")
                    cell = resolve_slot(w, roles, slot)
                    used_slots.add(slot.slot_id)
                    role, record_id, field = (binding[k] for k in ("record_role", "record_id", "field_name"))
                    profile = visual["display_profile_ref"]
                    token_id = None
                    if profile == "SIGNIFICANCE_TOKEN":
                        require(role == "comparison" and field == "status", "Token mapping mismatch")
                        candidates = [t for k, t in token_map.items() if k[0] == run.result_run_id
                            and t["token_id"] == slot.significance_token_id
                            and t["comparison_id"] == record_id
                            and t["envelope_checksum"] in visual["significance_presentation_refs"]
                            and (slot.significance_value_id is None or t["value_id"] == slot.significance_value_id)]
                        require(len(candidates) == 1, "Unknown/ambiguous significance token")
                        token = candidates[0]
                        require(token["comparison_id"] == record_id, "Token source mapping mismatch")
                        require(token["envelope_checksum"] in visual["significance_presentation_refs"],
                                "Token envelope not pinned by visual")
                        token_id = token["token_id"]
                        kind, value, fmt, status, unit = "literal", token["display_text"], "General", "SIGNIFICANT", "TEXT"
                    else:
                        require(role in COLLECTIONS or role == "provenance", "Unknown source role")
                        if role == "provenance":
                            require(record_id == run.result_run_id, "Provenance source mismatch")
                            record = to_canonical_data(run)
                        else:
                            records = to_canonical_data(run)[COLLECTIONS[role]]
                            matches = [r for r in records if r[ID_FIELDS[role]] == record_id]
                            require(len(matches) == 1, "Missing/ambiguous canonical source")
                            record = matches[0]
                        require(field in record, "Unknown mapped source field")
                        kind, value, fmt, status, unit = present(record[field], record, field, profile, slot.storage_mode)
                    if slot.storage_mode == "exact_text":
                        fmt = cell.number_format
                    elif kind == "blank" and profile == "COUNT":
                        require(cell.number_format in ("0", "0.00"), "Unknown COUNT display format")
                        fmt = cell.number_format
                    require((slot.kind == "numeric" and kind in ("numeric", "blank"))
                            or (slot.kind == "literal" and kind in ("literal", "blank")), "Target type mismatch")
                    require(slot.status_name is not None, "Missing required status sidecar")
                    put(cell, kind, value, run, role, record_id, field, slot.slot_id,
                        visual["visual_id"], profile, slot.storage_mode, status, unit, token_id, fmt)
                    status_cell = named_cell(w, slot.status_name, cell.parent.title)
                    annotation = status
                    if slot.storage_mode == "exact_text" and profile not in ("LITERAL", "SIGNIFICANCE_TOKEN"):
                        annotation += " | EXACT_TEXT"
                    status_field = ("value_status" if role == "value" else "base_status" if role == "base"
                                    else "status" if role == "comparison" else "renderer_status")
                    put(status_cell, "literal", annotation, run, role, record_id, status_field, slot.slot_id + ":status",
                        visual["visual_id"], "LITERAL", number_format=status_cell.number_format)
                    mapped = canonical_payload(writes[(cell.parent.title, cell.coordinate)])
                    mapped["status_sidecar"] = canonical_payload(writes[(status_cell.parent.title, status_cell.coordinate)])
                    if profile not in ("LITERAL", "SIGNIFICANCE_TOKEN"):
                        require(slot.label_name is not None, "Missing declared label target")
                        label_cell = named_cell(w, slot.label_name, cell.parent.title)
                        labels = visual["labels"]
                        require(not labels or len(labels) == len(visual["source_bindings"]), "Ambiguous label mapping")
                        if labels:
                            caption = labels[binding_index]["text"]
                            label_role, label_field = "visual_spec", "labels"
                        else:
                            caption = record_id
                            label_role, label_field = role, ID_FIELDS.get(role, "result_run_id")
                            renderer_warnings.append({"code": "TECHNICAL_ID_LABEL", "visual_id": visual["visual_id"],
                                                      "record_id": record_id, "slot_id": slot.slot_id})
                        put(label_cell, "literal", caption, run, label_role, record_id, label_field,
                            slot.slot_id + ":label", visual["visual_id"], "LITERAL", number_format=label_cell.number_format)
                        mapped["label_target"] = canonical_payload(writes[(label_cell.parent.title, label_cell.coordinate)])
                    presentation.append(mapped)
        require(used_slots == set(slots), "Incomplete Render Plan / unused mapping")
        payloads = tuple(to_canonical_json(r) for r in runs.values())
        provenance = {"project_id": next(iter(runs.values())).project_id,
                      "configuration_id": request.configuration_id,
                      "configuration_version": request.configuration_version,
                      "configuration_checksum": config_sha, "configuration_snapshot": canonical_payload(spec),
                      "visual_spec_json": request.visual_spec_json,
                      "configuration_release": canonical_payload(request.configuration_release),
                      "master_release": canonical_payload(request.master_release),
                      "master": canonical_payload(request.master), "visual_spec_checksum": sha_bytes(request.visual_spec_json),
                      "renderer": BUILD, "mapping": MAPPING, "numeric": NUMERIC, "significance": SIGNIFICANCE,
                      "runs": [{k: to_canonical_data(r)[k] for k in
                                ("project_spec_ref", "dataset_fingerprint", "result_run_id", "result_fingerprint",
                                 "result_schema_version", "core_version", "manifest", "request", "release", "qa",
                                 "supersedes_result_run_id")}
                               for r in runs.values()],
                      "tokens": [decode(t) for t in request.significance_json],
                      "significance_json": request.significance_json,
                      "significance_letters": "SUPPLIED" if token_map else "LETTER_DISPLAY_UNSUPPORTED",
                      "warnings": [canonical_payload(q) for r in runs.values() for q in r.qa_events],
                      "renderer_warnings": renderer_warnings,
                      "errors": [], "vba_before": inventory["vba_sha256"],
                      "optional_source_versions": "UNAVAILABLE_FROM_CANONICAL_RESULT",
                      "audit_storage_mode": "EXACT_SERIALIZED_JSON_TEXT"}
        provenance["render_outcome_policy"] = "PUBLISH_ONLY_AFTER_PASS"
        audit = {"results": [], "bases": [], "slices": [], "significance": [], "qa": [],
                 "cell_map": [], "canonical_payload": [], "provenance": []}
        for run, payload in zip(runs.values(), payloads):
            data = to_canonical_data(run)
            for role, collection, id_field in (("results", "values", "value_id"), ("bases", "bases", "base_id"),
                    ("slices", "slices", "slice_id"), ("significance", "comparisons", "comparison_id")):
                for record in data[collection]:
                    audit[role].append((run, {"result_run_id": run.result_run_id,
                                            "result_fingerprint": run.result_fingerprint, **record}))
            qa_records = [(q["qa_id"], q) for q in data["qa_events"]]
            qa_records += [(i["issue_id"], i) for i in data["qa"]["issues"]]
            for record_id, record in qa_records:
                audit["qa"].append((run, {"record_id": record_id, "result_run_id": run.result_run_id,
                                          "payload": record}))
            for i in range(0, len(payload), 16000):
                audit["canonical_payload"].append((run, {"record_id": f"{run.result_run_id}:{i // 16000}",
                    "result_run_id": run.result_run_id, "chunk_id": i // 16000, "encoding": "UTF-8",
                    "total_length": len(payload), "checksum": sha_bytes(payload), "payload": payload[i:i+16000]}))
        for item in presentation:
            run = runs[item["result_run_id"]]
            audit["cell_map"].append((run, {"record_id": item["slot_id"], "result_run_id": run.result_run_id,
                                          "payload": item}))
        audit["provenance"].append((None, {"record_id": "render_inputs", "result_run_id": "ALL_RUNS",
                                           "payload": provenance}))
        for role, records in audit.items():
            binding = roles[role]
            sheet = w[binding.sheet]
            table = sheet.tables[binding.table]
            columns = [c.name for c in table.tableColumns]
            require(len(columns) == len(set(columns)), "Ambiguous audit columns")
            expected = audit_columns(role)
            require(set(columns) == set(expected), "Incompatible audit table schema")
            require(len(records) <= binding.capacity, "Audit capacity exceeded")
            left, top, _, _ = range_boundaries(table.ref)
            for index in range(binding.capacity):
                run, record = records[index] if index < len(records) else (None, None)
                for offset, column in enumerate(columns):
                    cell = sheet.cell(top + 1 + index, left + offset)
                    # Audit is explicitly exact serialized text, not approximate Excel numeric authority.
                    value = stable_json(record[column]) if record is not None else None
                    put(cell, "literal" if record is not None else "blank", value, run,
                        role, str(record.get(expected[0], index)) if record else f"stale:{index}",
                        column, f"{role}:{index}:{column}", number_format=cell.number_format)
        ordered = tuple(writes.values())
        require(sha(master_path) == request.master.artifact_sha256, "Master changed during planning")
        provenance_json = stable_json(provenance)
        plan_hash = fingerprint({"writes": ordered, "input": inventory["whole_sha256"],
                                 "canonical": payloads, "provenance": provenance_json, "request": config_sha})
        return RenderPlan(ordered, inventory["whole_sha256"], inventory["vba_sha256"], payloads,
                          provenance_json, config_sha, plan_hash)
    finally:
        w.close()
        w.vba_archive.close()


def sha_bytes(value):
    return hashlib.sha256(value.encode("utf-8")).hexdigest()


def audit_columns(role):
    model = {"results": CanonicalValue, "bases": CanonicalBase, "slices": CanonicalSlice,
             "significance": SignificanceRelation}.get(role)
    if model:
        return ("result_run_id", "result_fingerprint", *(f.name for f in fields(model)))
    if role == "canonical_payload":
        return ("record_id", "result_run_id", "chunk_id", "encoding", "total_length", "checksum", "payload")
    return ("record_id", "result_run_id", "payload")


def worksheet_parts(z):
    parser = etree.XMLParser(resolve_entities=False, no_network=True)
    book = etree.fromstring(z.read("xl/workbook.xml"), parser)
    rels = etree.fromstring(z.read("xl/_rels/workbook.xml.rels"), parser)
    result = {}
    for sheet in book.findall(f"{{{S}}}sheets/{{{S}}}sheet"):
        links = [r for r in rels if r.get("Id") == sheet.get(f"{{{R}}}id")]
        require(len(links) == 1 and links[0].get("TargetMode") != "External", "Invalid worksheet relationship")
        target = links[0].get("Target", "")
        part = target.lstrip("/") if target.startswith("/") else posixpath.normpath("xl/" + target)
        require(part in z.namelist() and sheet.get("name") not in result, "Invalid worksheet binding")
        result[sheet.get("name")] = part
    return result


def mutate(source, staged, plan):
    parser = etree.XMLParser(resolve_entities=False, no_network=True)
    with ZipFile(source) as original, ZipFile(staged, "w") as output:
        parts = worksheet_parts(original)
        edits = {}
        for write in plan.writes:
            require(write.sheet in parts, "Missing physical target")
            edits.setdefault(parts[write.sheet], []).append(write)
        for item in original.infolist():
            if item.filename not in edits:
                with original.open(item) as src, output.open(item, "w") as dst:
                    shutil.copyfileobj(src, dst, 1024 * 1024)
                continue
            tree = etree.fromstring(original.read(item.filename), parser)
            cells = {}
            for cell in tree.findall(f"{{{S}}}sheetData/{{{S}}}row/{{{S}}}c"):
                require(cell.get("r") not in cells, "Ambiguous physical cell")
                cells[cell.get("r")] = cell
            for write in edits[item.filename]:
                require(write.cell in cells, "Undeclared/unmaterialized physical target")
                cell = cells[write.cell]
                require(cell.find(f"{{{S}}}f") is None, "Protected formula overwrite")
                for child in list(cell):
                    cell.remove(child)
                if write.kind == "literal":
                    cell.set("t", "inlineStr")
                    inline = etree.SubElement(cell, f"{{{S}}}is")
                    txt = etree.SubElement(inline, f"{{{S}}}t")
                    txt.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
                    txt.text = literal(write.value)
                elif write.kind == "numeric":
                    cell.set("t", "n")
                    etree.SubElement(cell, f"{{{S}}}v").text = stable_json(write.value)
                elif write.kind == "blank":
                    cell.attrib.pop("t", None)
                else:
                    raise RenderError("Undeclared write type")
            output.writestr(item, etree.tostring(tree, xml_declaration=True, encoding="UTF-8", standalone=True))


def validate_output(source, output, plan):
    expected_hash = fingerprint({"writes": plan.writes, "input": plan.input_sha256,
                                 "canonical": plan.canonical_json, "provenance": plan.provenance_json,
                                 "request": plan.request_sha256})
    require(plan.plan_sha256 == expected_hash, "Render Plan integrity failure")
    require(len(plan.writes) == len({(p.sheet, p.cell) for p in plan.writes}) and bool(plan.writes),
            "Incomplete/ambiguous Render Plan")
    provenance = decode(plan.provenance_json)
    require(all(k in provenance for k in ("master", "runs", "configuration_checksum", "renderer",
                "mapping", "numeric", "significance", "warnings", "errors", "vba_before")), "Incomplete provenance")
    require(sha(source) == plan.input_sha256, "Master changed during rendering")
    inventory = inspect_package(Path(output))
    require(inventory["vba_sha256"] == plan.vba_sha256, "VBA preservation failure")
    parser = etree.XMLParser(resolve_entities=False, no_network=True)
    with ZipFile(source) as a, ZipFile(output) as b:
        require(set(a.namelist()) == set(b.namelist()), "Package inventory changed")
        parts = worksheet_parts(a)
        edited = {}
        for write in plan.writes:
            edited.setdefault(parts[write.sheet], set()).add(write.cell)
        for name in a.namelist():
            if name not in edited:
                require(hashlib.sha256(a.read(name)).digest() == hashlib.sha256(b.read(name)).digest(),
                        "Untouched package part changed")
                continue
            trees = [etree.fromstring(z.read(name), parser) for z in (a, b)]
            for tree in trees:
                removed = set()
                for cell in tree.findall(f"{{{S}}}sheetData/{{{S}}}row/{{{S}}}c"):
                    if cell.get("r") in edited[name]:
                        require(cell.get("r") not in removed, "Duplicate output cell")
                        removed.add(cell.get("r"))
                        # Style and target identity remain immutable even for writable content.
                        for child in list(cell):
                            cell.remove(child)
                        cell.attrib.pop("t", None)
                require(removed == edited[name], "Missing output target")
            require(etree.tostring(trees[0], method="c14n") == etree.tostring(trees[1], method="c14n"),
                    "Undeclared worksheet mutation")
    w = load_workbook(output, keep_vba=True, keep_links=True)
    try:
        for write in plan.writes:
            cell = w[write.sheet][write.cell]
            require(stable_json(cell.value) == stable_json(write.value) and type(cell.value) is type(write.value),
                    "Exact scalar/type readback failed")
            require(cell.number_format == write.number_format, "Display format changed")
            require(cell.data_type != "f" and (write.kind != "literal" or cell.data_type == "s"),
                    "Literal/formula safety failed")
    finally:
        w.close()
        w.vba_archive.close()
    return {"qa": "PASS", "vba_before": plan.vba_sha256, "vba_after": inventory["vba_sha256"],
            "output_sha256": sha(output), "actual_targets": [(p.sheet, p.cell) for p in plan.writes],
            "untouched_parts_exact": True, "undeclared_cells_preserved": True,
            "plan": canonical_payload(plan), "warnings": decode(plan.provenance_json)["warnings"], "errors": []}


def render(request, master_path, output_path, *, plan=None):
    master_path, output_path = Path(master_path), Path(output_path)
    require(output_path.suffix.lower() == ".xlsm" and not output_path.exists()
            and output_path.resolve() != master_path.resolve(), "Invalid output destination")
    require(output_path.parent.is_dir(), "Output directory absent")
    # The stage is on the output filesystem; no workbook is published until all QA passes.
    with TemporaryDirectory(prefix="m7b_stage_", dir=output_path.parent) as temp:
        snapshot = Path(temp) / "master_snapshot.xlsm"
        shutil.copyfile(master_path, snapshot)
        require(sha(snapshot) == request.master.artifact_sha256, "Master snapshot identity mismatch")
        expected = plan_render(request, snapshot)
        require(plan is None or plan == expected, "Incomplete/tampered Render Plan")
        staged = Path(temp) / "staged.xlsm"
        mutate(snapshot, staged, expected)
        evidence = validate_output(snapshot, staged, expected)
        require(sha(master_path) == expected.input_sha256, "Input Master changed before publication")
        require(not output_path.exists(), "Output destination appeared during rendering")
        # Hard-link publication refuses a racing existing destination instead of overwriting it.
        import os
        os.link(staged, output_path)
    return evidence
