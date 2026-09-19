from __future__ import annotations

from copy import copy
from dataclasses import dataclass
import hashlib
from pathlib import Path
import shutil
from tempfile import TemporaryDirectory
from typing import Any
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl.formula.translate import Translator
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import range_boundaries
from openpyxl.worksheet.table import TableColumn

from src.analytics_core.result_identity import canonical_payload, fingerprint
from src.analytics_core.serialization import to_canonical_data
from src.contracts.models import CanonicalResult

from .renderer import (
    RenderError, Slot, literal, present, require, sha, sha_bytes, tokens, validate_sources,
)

DYNAMIC_CONTRACT = "M7_DYNAMIC_REGION_CONTRACT_V1"
DYNAMIC_BINDING_SCHEMA = "M7_DYNAMIC_REGION_BINDING_V1"
DYNAMIC_QUALIFICATION_SCHEMA = "M7_DYNAMIC_REGION_QUALIFICATION_V1"
DYNAMIC_PLAN = "M7_DYNAMIC_REGION_RENDER_PLAN_V1"
DYNAMIC_BUILD = "M7_DYNAMIC_REGION_RUNTIME_V1"
COLLISION_POLICY = "M7_DYNAMIC_REGION_COLLISION_V1"


@dataclass(frozen=True)
class Bounds:
    left: int
    top: int
    right: int
    bottom: int

    @property
    def rows(self):
        return self.bottom - self.top + 1

    @property
    def columns(self):
        return self.right - self.left + 1

    def cells(self):
        return tuple((r, c) for r in range(self.top, self.bottom + 1)
                     for c in range(self.left, self.right + 1))

    def intersects(self, other):
        return not (self.right < other.left or other.right < self.left
                    or self.bottom < other.top or other.bottom < self.top)


@dataclass(frozen=True)
class DynamicRegionStylePolicy:
    style_policy_id: str
    region_id: str
    template_cell: str
    propagation_axis: str
    owned_attributes: tuple[str, ...]
    unsupported_attributes: tuple[str, ...] = ()
    table_style_policy: str = "NOT_APPLICABLE"


@dataclass(frozen=True)
class DynamicRegionFormulaPolicy:
    formula_policy_id: str
    region_id: str
    classification: str
    source_formula_cells: tuple[str, ...]
    target_field_offsets: tuple[int, ...]
    propagation_axis: str
    normalized_formula_rule: str = "OPENPYXL_TRANSLATOR_RELATIVE_V1"
    expected_inventory_rule: str = "EXACT_DECLARED_TARGETS_V1"


@dataclass(frozen=True)
class DynamicField:
    field_id: str
    source_field: str
    column_offset: int
    display_profile_ref: str
    storage_mode: str


@dataclass(frozen=True)
class DynamicRegionDeclaration:
    region_id: str
    role_id: str
    worksheet_role: str
    sheet: str
    region_kind: str
    table_id: str | None
    anchor_name: str | None
    growth_dimensions: str
    min_rows: int
    min_columns: int
    max_rows: int
    max_columns: int
    owned_envelope: Bounds
    field_columns: tuple[DynamicField, ...]
    style_policy_ref: str
    formula_policy_ref: str = "NONE"
    region_contract_version: str = DYNAMIC_CONTRACT
    ownership: str = "RENDERER_OWNED"
    cleanup_policy: str = "REPLACE_AND_CLEAR_STALE"
    collision_policy_ref: str = COLLISION_POLICY
    update_policy: str = "REPLACE"


@dataclass(frozen=True)
class DynamicRegionBinding:
    region_id: str
    result_run_id: str
    record_role: str
    ordered_record_ids: tuple[str, ...]
    field_bindings: tuple[str, ...]
    requested_rows: int
    requested_columns: int
    display_profile_refs: tuple[str, ...]
    significance_presentation_refs: tuple[str, ...] = ()


@dataclass(frozen=True)
class DynamicMasterQualification:
    schema_version: str
    master_id: str
    master_version: str
    master_sha256: str
    region_contract_version: str
    declarations: tuple[DynamicRegionDeclaration, ...]
    style_policies: tuple[DynamicRegionStylePolicy, ...]
    formula_policies: tuple[DynamicRegionFormulaPolicy, ...] = ()
    fixed_slots: tuple[Slot, ...] = ()


@dataclass(frozen=True)
class DynamicLineage:
    qualified_master_sha256: str
    prior_output_sha256: str
    region_bounds: tuple[tuple[str, Bounds], ...]
    prior_plan_sha256: str


@dataclass(frozen=True)
class DynamicRenderRequest:
    results: tuple[CanonicalResult, ...]
    qualification: DynamicMasterQualification
    binding_schema_version: str
    bindings: tuple[DynamicRegionBinding, ...]
    significance_json: tuple[str, ...] = ()
    lineage: DynamicLineage | None = None
    renderer_build: str = DYNAMIC_BUILD


@dataclass(frozen=True)
class DynamicRegionOperation:
    operation_id: str
    operation_type: str
    region_id: str
    result_run_id: str
    source_record_ids: tuple[str, ...]
    before_bounds: Bounds
    after_bounds: Bounds
    requested_rows: int
    requested_columns: int
    owned_envelope: Bounds
    affected_cells: tuple[str, ...]
    policy_refs: tuple[str, ...]
    collision_validation: str


@dataclass(frozen=True)
class DynamicWrite:
    region_id: str
    sheet: str
    cell: str
    kind: str
    value: str | int | float | None
    number_format: str
    record_id: str
    field_id: str
    token_id: str | None = None
    comparison_id: str | None = None


@dataclass(frozen=True)
class DynamicRenderPlan:
    plan_version: str
    contract_version: str
    input_sha256: str
    qualified_master_sha256: str
    declaration_sha256: str
    binding_sha256: str
    dynamic_operations: tuple[DynamicRegionOperation, ...]
    writes: tuple[DynamicWrite, ...]
    vba_sha256: str
    protected_parts_sha256: str
    plan_sha256: str


def _vba(path):
    with ZipFile(path) as package:
        return hashlib.sha256(package.read("xl/vbaProject.bin")).hexdigest()


def _protected(path, declarations):
    excluded = {(d.sheet, row, col) for d in declarations for row, col in d.owned_envelope.cells()}
    region_tables = {d.table_id for d in declarations if d.table_id}
    region_anchors = {d.anchor_name for d in declarations if d.anchor_name}
    w = load_workbook(path, keep_vba=True, keep_links=True, data_only=False)
    try:
        cells = []
        tables = []
        for sheet in w.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    if (sheet.title, cell.row, cell.column) not in excluded and cell.value is not None:
                        cells.append((sheet.title, cell.coordinate, cell.value, cell.data_type,
                                      cell.number_format, cell.style_id, cell.protection.locked))
            for table in sheet.tables.values():
                if table.name not in region_tables:
                    tables.append((sheet.title, table.name, table.ref, table.tableStyleInfo.name if table.tableStyleInfo else None))
        names = sorted((name, defined.attr_text) for name, defined in w.defined_names.items()
                       if name not in region_anchors)
        merges = sorted((sheet.title, str(item)) for sheet in w.worksheets for item in sheet.merged_cells.ranges)
    finally:
        w.close()
        if w.vba_archive:
            w.vba_archive.close()
    with ZipFile(path) as package:
        sensitive_prefixes = ("xl/charts/", "xl/drawings/", "xl/pivot", "xl/slicers/",
                              "xl/ctrlProps/", "xl/activeX/", "xl/externalLinks/")
        parts = sorted((name, hashlib.sha256(package.read(name)).hexdigest())
                       for name in package.namelist()
                       if name == "xl/vbaProject.bin" or name.startswith(sensitive_prefixes))
    return fingerprint({"sheets": tuple(w.sheetnames), "cells": cells, "tables": tables,
                        "names": names, "merges": merges, "sensitive_parts": parts})


def _address(bounds):
    return f"{get_column_letter(bounds.left)}{bounds.top}:{get_column_letter(bounds.right)}{bounds.bottom}"


def _cells(sheet, bounds):
    return tuple(f"{sheet}!{get_column_letter(c)}{r}" for r, c in bounds.cells())


def _policy_fingerprint(value):
    return fingerprint(canonical_payload(value))


def _validate_contract(request, path):
    q = request.qualification
    require(request.renderer_build == DYNAMIC_BUILD, "Unsupported Dynamic Region build")
    require(request.binding_schema_version == DYNAMIC_BINDING_SCHEMA, "Wrong binding contract version")
    require(q.schema_version == DYNAMIC_QUALIFICATION_SCHEMA
            and q.region_contract_version == DYNAMIC_CONTRACT, "Wrong Dynamic Region contract version")
    require(sha(path) == (request.lineage.prior_output_sha256 if request.lineage else q.master_sha256),
            "Wrong Master hash / unverifiable previous output")
    if request.lineage:
        require(request.lineage.qualified_master_sha256 == q.master_sha256
                and bool(request.lineage.prior_plan_sha256), "Unverifiable previous active footprint")
    ids = [d.region_id for d in q.declarations]
    require(ids and len(ids) == len(set(ids)), "Duplicate region identity")
    require(len(request.bindings) == len({b.region_id for b in request.bindings}), "Duplicate region binding")
    require({b.region_id for b in request.bindings} <= set(ids), "Unknown region_id")
    return validate_sources(request)


def _resolve_anchor(w, declaration):
    if declaration.region_kind == "TABLE":
        require(declaration.table_id and not declaration.anchor_name, "Malformed table binding")
        matches = [(s, s.tables[declaration.table_id]) for s in w.worksheets
                   if declaration.table_id in s.tables]
        require(len(matches) == 1 and matches[0][0].title == declaration.sheet, "Missing/wrong table identity")
        sheet, table = matches[0]
        left, top, right, bottom = range_boundaries(table.ref)
        return sheet, table, Bounds(left, top + 1, right, bottom)
    require(declaration.region_kind == "RANGE" and declaration.anchor_name and not declaration.table_id,
            "Malformed range binding")
    require(declaration.anchor_name in w.defined_names, "Missing named anchor")
    destinations = list(w.defined_names[declaration.anchor_name].destinations)
    require(len(destinations) == 1, "Ambiguous named anchor")
    sheet_name, address = destinations[0]
    left, top, right, bottom = range_boundaries(address.replace("$", ""))
    require(sheet_name == declaration.sheet and left == right and top == bottom, "Wrong named anchor identity")
    return w[sheet_name], None, Bounds(left, top, left + declaration.min_columns - 1,
                                       top + declaration.min_rows - 1)


def _validate_declaration(d, style, formula):
    require(d.region_contract_version == DYNAMIC_CONTRACT and d.region_id and d.role_id
            and d.worksheet_role and d.sheet, "Invalid region identity")
    require(d.growth_dimensions in ("ROWS", "COLUMNS", "ROWS_AND_COLUMNS"), "Unknown growth dimension")
    require(all(type(v) is int and v > 0 for v in (d.min_rows, d.min_columns, d.max_rows, d.max_columns))
            and d.min_rows <= d.max_rows and d.min_columns <= d.max_columns, "Invalid finite envelope")
    require(d.owned_envelope.rows == d.max_rows and d.owned_envelope.columns == d.max_columns,
            "Owned envelope mismatch")
    require(d.ownership == "RENDERER_OWNED" and d.cleanup_policy == "REPLACE_AND_CLEAR_STALE"
            and d.update_policy == "REPLACE" and d.collision_policy_ref == COLLISION_POLICY,
            "Unsupported ownership/update policy")
    require(style and style.region_id == d.region_id and style.propagation_axis == d.growth_dimensions,
            "Malformed style policy")
    allowed_attrs = {"font", "fill", "border", "alignment", "number_format", "protection",
                     "row_height", "column_width", "conditional_formatting"}
    require(bool(style.owned_attributes) and set(style.owned_attributes) <= allowed_attrs
            and set(style.unsupported_attributes) <= allowed_attrs
            and not set(style.owned_attributes) & set(style.unsupported_attributes),
            "Malformed style policy")
    require("conditional_formatting" not in style.owned_attributes,
            "Unsupported conditional formatting propagation")
    if d.region_kind == "TABLE":
        require(style.table_style_policy == "PRESERVE_EXISTING_TABLE_STYLE", "Invalid TABLE style policy")
    else:
        require(style.table_style_policy == "NOT_APPLICABLE", "Invalid RANGE style policy")
    if d.formula_policy_ref == "NONE":
        require(formula is None, "Contradictory formula policy")
    else:
        require(formula and formula.region_id == d.region_id
                and formula.classification == "PRESENTATION_ONLY"
                and formula.normalized_formula_rule == "OPENPYXL_TRANSLATOR_RELATIVE_V1"
                and formula.expected_inventory_rule == "EXACT_DECLARED_TARGETS_V1",
                "Analytical or malformed formula policy")
        require(len(formula.source_formula_cells) == 1
                and formula.propagation_axis == d.growth_dimensions
                and formula.target_field_offsets
                and all(type(offset) is int and offset >= 0 for offset in formula.target_field_offsets)
                and len(set(formula.target_field_offsets)) == len(formula.target_field_offsets),
                "Malformed formula source/target policy")


def _fixed_slot_cell(w, slot):
    require(bool(slot.name) != bool(slot.table), "Ambiguous/unresolvable fixed slot")
    if slot.name:
        require(slot.name in w.defined_names and slot.row is None and slot.column is None,
                "Ambiguous/unresolvable fixed slot")
        destinations = list(w.defined_names[slot.name].destinations)
        require(len(destinations) == 1, "Ambiguous/unresolvable fixed slot")
        sheet, address = destinations[0]
        left, top, right, bottom = range_boundaries(address.replace("$", ""))
        require(left == right and top == bottom, "Ambiguous/unresolvable fixed slot")
        return sheet, top, left
    require(type(slot.row) is int and slot.row >= 0 and isinstance(slot.column, str),
            "Ambiguous/unresolvable fixed slot")
    matches = [(sheet, sheet.tables[slot.table]) for sheet in w.worksheets if slot.table in sheet.tables]
    require(len(matches) == 1, "Ambiguous/unresolvable fixed slot")
    sheet, table = matches[0]
    columns = [column.name for column in table.tableColumns]
    require(columns.count(slot.column) == 1, "Ambiguous/unresolvable fixed slot")
    left, top, _, bottom = range_boundaries(table.ref)
    require(top + 1 + slot.row <= bottom, "Ambiguous/unresolvable fixed slot")
    return sheet.title, top + 1 + slot.row, left + columns.index(slot.column)


def _dynamic_tokens(request, runs):
    refs = {ref for binding in request.bindings for ref in binding.significance_presentation_refs}
    fake_spec = {"sections": [{"visuals": [{"significance_presentation_refs": list(refs)}]}]}
    class TokenRequest:
        significance_json = request.significance_json
    token_map = tokens(TokenRequest(), runs, fake_spec)
    require(bool(request.significance_json) == bool(refs), "Unused or unresolved significance envelope")
    return token_map


def plan_dynamic_render(request: DynamicRenderRequest, input_path: Path) -> DynamicRenderPlan:
    input_path = Path(input_path)
    runs = _validate_contract(request, input_path)
    q = request.qualification
    token_map = _dynamic_tokens(request, runs)
    styles = {p.style_policy_id: p for p in q.style_policies}
    formulas = {p.formula_policy_id: p for p in q.formula_policies}
    require(len(styles) == len(q.style_policies) and len(formulas) == len(q.formula_policies),
            "Duplicate policy identity")
    prior = dict(request.lineage.region_bounds) if request.lineage else {}
    require(not request.lineage or set(prior) == {d.region_id for d in q.declarations},
            "Unverifiable previous active footprint")
    w = load_workbook(input_path, keep_vba=True, keep_links=True)
    operations, writes, projected = [], [], {}
    consumed_significance_refs = set()
    try:
        declarations = {d.region_id: d for d in q.declarations}
        bindings = {b.region_id: b for b in request.bindings}
        for d in q.declarations:
            style, formula = styles.get(d.style_policy_ref), formulas.get(d.formula_policy_ref)
            _validate_declaration(d, style, formula)
            sheet, table, physical_before = _resolve_anchor(w, d)
            template_sheet, template_address = style.template_cell.split("!", 1) if style else ("", "")
            require(template_sheet == d.sheet, "Unknown style source")
            before = prior.get(d.region_id, physical_before)
            require(before.left == d.owned_envelope.left and before.top == d.owned_envelope.top
                    and before.right <= d.owned_envelope.right and before.bottom <= d.owned_envelope.bottom,
                    "Unverifiable previous active footprint")
            b = bindings.get(d.region_id)
            require(b is not None, "Missing Dynamic Region binding")
            require(len(b.significance_presentation_refs) == len(set(b.significance_presentation_refs)),
                    "Duplicate significance envelope reference")
            require(type(b.requested_rows) is int and type(b.requested_columns) is int
                    and b.requested_rows > 0 and b.requested_columns > 0, "Invalid requested extent")
            require(d.min_rows <= b.requested_rows <= d.max_rows
                    and d.min_columns <= b.requested_columns <= d.max_columns, "Dynamic Region overflow")
            require(d.growth_dimensions in ("ROWS", "ROWS_AND_COLUMNS") or b.requested_rows == d.min_rows,
                    "Growth in unauthorized row dimension")
            require(d.growth_dimensions in ("COLUMNS", "ROWS_AND_COLUMNS") or b.requested_columns == d.min_columns,
                    "Growth in unauthorized column dimension")
            require(len(b.ordered_record_ids) == b.requested_rows and b.field_bindings
                    and len(b.field_bindings) == b.requested_columns
                    and len(b.display_profile_refs) == len(b.field_bindings), "Cardinality mismatch")
            fields = {f.field_id: f for f in d.field_columns}
            require(len(fields) == len(d.field_columns) and tuple(b.field_bindings) == tuple(f.field_id for f in d.field_columns[:b.requested_columns]),
                    "Unknown or implicit field selection")
            require(tuple(b.display_profile_refs) == tuple(fields[x].display_profile_ref for x in b.field_bindings),
                    "Display profile mismatch")
            run = runs.get(b.result_run_id)
            require(run is not None, "Missing canonical record")
            role_collections = {"value": ("values", "value_id"), "base": ("bases", "base_id"),
                                "slice": ("slices", "slice_id"), "comparison": ("comparisons", "comparison_id"),
                                "qa": ("qa_events", "qa_id")}
            if b.record_role == "provenance":
                allowed = {"project_id", "dataset_fingerprint", "project_spec_ref", "core_version",
                           "result_run_id", "result_fingerprint", "result_schema_version",
                           "supersedes_result_run_id"}
                record = {key: value for key, value in to_canonical_data(run).items() if key in allowed}
                records = {run.result_run_id: record}
                id_field = "result_run_id"
            else:
                require(b.record_role in role_collections, "Wildcard/implicit source selection prohibited")
                collection, id_field = role_collections[b.record_role]
                records = {r[id_field]: r for r in to_canonical_data(run)[collection]}
            require(len(b.ordered_record_ids) == len(set(b.ordered_record_ids)), "Duplicate canonical record")
            require(set(b.ordered_record_ids) <= set(records), "Missing canonical record")
            after = Bounds(d.owned_envelope.left, d.owned_envelope.top,
                           d.owned_envelope.left + b.requested_columns - 1,
                           d.owned_envelope.top + b.requested_rows - 1)
            projected[d.region_id] = after
            # Every envelope is exclusive in V1, even when active footprints do not overlap.
            for other in q.declarations:
                require(other.region_id == d.region_id or other.sheet != d.sheet
                        or not d.owned_envelope.intersects(other.owned_envelope),
                        "Dynamic Region overlap")
            # Fixed targets, unrelated tables, merged cells and defined names are protected.
            for slot in q.fixed_slots:
                fixed_sheet, fixed_row, fixed_col = _fixed_slot_cell(w, slot)
                require(fixed_sheet != d.sheet or not d.owned_envelope.intersects(
                    Bounds(fixed_col, fixed_row, fixed_col, fixed_row)),
                    "Dynamic Region vs fixed-slot collision")
            for ws in w.worksheets:
                for candidate in ws.tables.values():
                    if ws.title == d.sheet and candidate.name == d.table_id:
                        continue
                    l, t, r, bot = range_boundaries(candidate.ref)
                    require(ws.title != d.sheet or not d.owned_envelope.intersects(Bounds(l, t, r, bot)),
                            "Unrelated Excel Table collision")
            for merged in sheet.merged_cells.ranges:
                l, t, r, bot = range_boundaries(str(merged))
                require(not d.owned_envelope.intersects(Bounds(l, t, r, bot)), "Merged-cell collision")
            for name, defined in w.defined_names.items():
                if name == d.anchor_name:
                    continue
                for sn, addr in defined.destinations:
                    l, t, r, bot = range_boundaries(addr.replace("$", ""))
                    require(sn != d.sheet or not d.owned_envelope.intersects(Bounds(l, t, r, bot)),
                            "Protected named range collision")
            allowed_formula_sources = set(formula.source_formula_cells if formula else ())
            mutable = set(before.cells()) | set(after.cells())
            for row, col in mutable:
                cell = sheet.cell(row, col)
                address = f"{sheet.title}!{cell.coordinate}"
                if cell.data_type == "f":
                    require(address in allowed_formula_sources, "Unallowlisted formula collision")
                elif (cell.value is not None and address != style.template_cell
                      and (request.lineage is None or (row, col) not in before.cells())):
                    require(False, "Protected/non-owned populated-cell collision")
            require(style.template_cell in {f"{sheet.title}!{sheet.cell(r, c).coordinate}" for r, c in d.owned_envelope.cells()},
                    "Unknown style source")
            if formula:
                source_sheet, source_address = formula.source_formula_cells[0].split("!", 1)
                source_left, source_top, source_right, source_bottom = range_boundaries(source_address)
                require(source_sheet == d.sheet and source_left == source_right and source_top == source_bottom
                        and d.owned_envelope.intersects(Bounds(source_left, source_top, source_right, source_bottom)),
                        "Invalid formula source identity")
                source_cell = sheet[source_address]
                require(source_cell.data_type == "f" and isinstance(source_cell.value, str)
                        and source_cell.value.startswith("="), "Formula source is not an Excel formula")
                formula_fields = {f.column_offset for f in d.field_columns
                                  if f.source_field == "__PRESENTATION_FORMULA__"}
                require(set(formula.target_field_offsets) == formula_fields
                        and all(offset < d.max_columns for offset in formula.target_field_offsets),
                        "Formula targets do not map to declared presentation fields")
            policy_refs = (DYNAMIC_CONTRACT, d.style_policy_ref, d.formula_policy_ref, COLLISION_POLICY)
            collision_hash = fingerprint({"region": d.region_id, "before": before, "after": after,
                                          "envelope": d.owned_envelope, "policies": policy_refs})
            def op(kind, cells):
                payload = {"kind": kind, "region": d.region_id, "run": b.result_run_id,
                           "records": b.ordered_record_ids, "before": before, "after": after,
                           "cells": cells, "policies": policy_refs, "collision": collision_hash}
                return DynamicRegionOperation(fingerprint(payload), kind, d.region_id, b.result_run_id,
                    b.ordered_record_ids, before, after, b.requested_rows, b.requested_columns,
                    d.owned_envelope, cells, policy_refs, "PASS:" + collision_hash)
            operations.append(op("VALIDATE_REGION", _cells(d.sheet, after)))
            if d.region_kind == "TABLE" and before != after:
                operations.append(op("RESIZE_TABLE", (_address(before), _address(after))))
            stale = tuple(sorted(set(_cells(d.sheet, before)) - set(_cells(d.sheet, after))))
            if stale:
                operations.append(op("CLEAR_OWNED_STALE", stale))
            operations.append(op("PROPAGATE_STYLE", _cells(d.sheet, after)))
            if formula:
                targets = tuple(f"{d.sheet}!{get_column_letter(after.left + offset)}{row}"
                    for offset in formula.target_field_offsets if offset < b.requested_columns
                    for row in range(after.top, after.bottom + 1))
                if targets:
                    operations.append(op("PROPAGATE_PRESENTATION_FORMULA", targets))
            for row_offset, record_id in enumerate(b.ordered_record_ids):
                record = records[record_id]
                for field_id in b.field_bindings:
                    field = fields[field_id]
                    if field.source_field == "__PRESENTATION_FORMULA__":
                        require(formula is not None and field.column_offset in formula.target_field_offsets,
                                "Unclassified presentation formula field")
                        continue
                    if field.source_field == "__SIGNIFICANCE_TOKEN__":
                        require(b.record_role == "value" and b.significance_presentation_refs,
                                "Significance marker requires explicit value binding")
                        candidates = [token for key, token in token_map.items()
                            if key[0] == run.result_run_id and token["value_id"] == record_id
                            and token["envelope_checksum"] in b.significance_presentation_refs]
                        require(len(candidates) == 1, "Unknown/ambiguous Dynamic significance marker")
                        token = candidates[0]
                        consumed_significance_refs.add(token["envelope_checksum"])
                        cell = sheet.cell(after.top + row_offset, after.left + field.column_offset)
                        writes.append(DynamicWrite(d.region_id, d.sheet, cell.coordinate, "literal",
                            token["display_text"], "General", record_id, field_id,
                            token["token_id"], token["comparison_id"]))
                        continue
                    if b.record_role == "provenance":
                        require(field.source_field in {"project_id", "dataset_fingerprint", "project_spec_ref",
                            "core_version", "result_run_id", "result_fingerprint", "result_schema_version",
                            "supersedes_result_run_id"}, "Unknown provenance field")
                    require(field.source_field in record, "Unknown canonical field")
                    kind, value, number_format, _, _ = present(record[field.source_field], record,
                        field.source_field, field.display_profile_ref, field.storage_mode)
                    cell = sheet.cell(after.top + row_offset, after.left + field.column_offset)
                    require(cell.data_type != "f", "Formula target cannot be canonical write")
                    writes.append(DynamicWrite(d.region_id, d.sheet, cell.coordinate, kind, value,
                                               number_format, record_id, field_id))
            region_writes = tuple(f"{x.sheet}!{x.cell}" for x in writes if x.region_id == d.region_id)
            operations.append(op("WRITE_CELL", region_writes))
        require(consumed_significance_refs == {sha_bytes(raw) for raw in request.significance_json},
                "Unused Dynamic significance envelope")
        order = {k: i for i, k in enumerate(("VALIDATE_REGION", "RESIZE_TABLE", "CLEAR_OWNED_STALE",
            "PROPAGATE_STYLE", "PROPAGATE_PRESENTATION_FORMULA", "WRITE_CELL"))}
        operations.sort(key=lambda x: (order[x.operation_type], x.region_id, x.operation_id))
        require(len({x.operation_id for x in operations}) == len(operations), "Duplicate operation identity")
        declaration_sha = _policy_fingerprint((q.declarations, q.style_policies, q.formula_policies))
        binding_sha = _policy_fingerprint(request.bindings)
        payload = {"version": DYNAMIC_PLAN, "contract": DYNAMIC_CONTRACT, "input": sha(input_path),
                   "master": q.master_sha256, "declarations": declaration_sha, "bindings": binding_sha,
                   "operations": operations, "writes": writes, "build": request.renderer_build}
        plan_sha = fingerprint(payload)
        return DynamicRenderPlan(DYNAMIC_PLAN, DYNAMIC_CONTRACT, sha(input_path), q.master_sha256,
            declaration_sha, binding_sha, tuple(operations), tuple(writes), _vba(input_path),
            _protected(input_path, q.declarations), plan_sha)
    finally:
        w.close()
        if w.vba_archive:
            w.vba_archive.close()


def _apply_style(source, target, policy):
    for attr in policy.owned_attributes:
        if attr in ("font", "fill", "border", "alignment", "protection"):
            setattr(target, attr, copy(getattr(source, attr)))
        elif attr == "number_format":
            target.number_format = source.number_format


def _logical_state(path, declarations):
    w = load_workbook(path, keep_vba=True, data_only=False)
    try:
        state = {}
        for d in declarations:
            s = w[d.sheet]
            state[d.region_id] = tuple((s.cell(r, c).coordinate, s.cell(r, c).value,
                                        s.cell(r, c).number_format, s.cell(r, c).style_id)
                                       for r, c in d.owned_envelope.cells())
        return state
    finally:
        w.close()
        if w.vba_archive:
            w.vba_archive.close()


def _style_value(cell, attribute):
    value = getattr(cell, attribute)
    return str(value) if attribute in ("font", "fill", "border", "alignment", "protection") else value


def _validate_dynamic_output(source_path, output_path, request, plan):
    declarations = {d.region_id: d for d in request.qualification.declarations}
    styles = {p.style_policy_id: p for p in request.qualification.style_policies}
    formulas = {p.formula_policy_id: p for p in request.qualification.formula_policies}
    source = load_workbook(source_path, keep_vba=True, data_only=False)
    output = load_workbook(output_path, keep_vba=True, data_only=False)
    evidence = {"tables": [], "writes": [], "styles": [], "formulas": [], "stale": [], "unchanged": []}
    try:
        writes = {(write.sheet, write.cell): write for write in plan.writes}
        affected = set(writes)
        expected_formulas = {}
        for operation in plan.dynamic_operations:
            d = declarations[operation.region_id]
            sheet = output[d.sheet]
            if operation.operation_type == "RESIZE_TABLE":
                table = sheet.tables[d.table_id]
                expected_ref = (f"{get_column_letter(operation.after_bounds.left)}{operation.after_bounds.top - 1}:"
                                f"{get_column_letter(operation.after_bounds.right)}{operation.after_bounds.bottom}")
                require(table.name == d.table_id and table.ref == expected_ref,
                        "Dynamic TABLE structure readback failed")
                require([column.name for column in table.tableColumns]
                        == [field.field_id for field in d.field_columns[:operation.requested_columns]],
                        "Dynamic TABLE column identity readback failed")
                source_table = source[d.sheet].tables[d.table_id]
                require(str(table.tableStyleInfo) == str(source_table.tableStyleInfo),
                        "Dynamic TABLE style identity readback failed")
                evidence["tables"].append((d.region_id, table.ref, tuple(c.name for c in table.tableColumns)))
            elif operation.operation_type == "CLEAR_OWNED_STALE":
                for item in operation.affected_cells:
                    sheet_name, address = item.split("!", 1)
                    cell = output[sheet_name][address]
                    require(cell.value is None and cell.data_type != "f", "Stale area readback failed")
                    affected.add((sheet_name, address))
                    evidence["stale"].append(item)
            elif operation.operation_type == "PROPAGATE_STYLE":
                policy = styles[d.style_policy_ref]
                template = source[d.sheet][policy.template_cell.split("!", 1)[1]]
                for item in operation.affected_cells:
                    sheet_name, address = item.split("!", 1)
                    target = output[sheet_name][address]
                    for attribute in policy.owned_attributes:
                        if attribute == "row_height":
                            require(sheet.row_dimensions[target.row].height
                                    == source[d.sheet].row_dimensions[template.row].height,
                                    "Dynamic row-height readback failed")
                            continue
                        if attribute == "column_width":
                            require(sheet.column_dimensions[get_column_letter(target.column)].width
                                    == source[d.sheet].column_dimensions[get_column_letter(template.column)].width,
                                    "Dynamic column-width readback failed")
                            continue
                        if attribute == "number_format" and (sheet_name, address) in writes:
                            continue
                        require(_style_value(target, attribute) == _style_value(template, attribute),
                                "Dynamic style readback failed")
                    affected.add((sheet_name, address))
                    evidence["styles"].append(item)
            elif operation.operation_type == "PROPAGATE_PRESENTATION_FORMULA":
                policy = formulas[d.formula_policy_ref]
                origin = policy.source_formula_cells[0].split("!", 1)[1]
                formula = source[d.sheet][origin].value
                for item in operation.affected_cells:
                    sheet_name, address = item.split("!", 1)
                    expected_formulas[(sheet_name, address)] = Translator(formula, origin=origin).translate_formula(address)
                    affected.add((sheet_name, address))
        for key, write in writes.items():
            cell = output[write.sheet][write.cell]
            require(type(cell.value) is type(write.value) and cell.value == write.value,
                    "Dynamic canonical write value readback failed")
            require(cell.number_format == write.number_format, "Dynamic canonical write format readback failed")
            require((write.kind != "literal" or cell.data_type == "s")
                    and (write.kind != "numeric" or cell.data_type == "n")
                    and (write.kind != "blank" or cell.value is None), "Dynamic canonical write type readback failed")
            evidence["writes"].append((write.sheet, write.cell, write.record_id, write.field_id,
                                       write.token_id, write.comparison_id))
        actual_formulas = {}
        for d in declarations.values():
            for row, col in d.owned_envelope.cells():
                cell = output[d.sheet].cell(row, col)
                if cell.data_type == "f":
                    actual_formulas[(d.sheet, cell.coordinate)] = cell.value
        # Controlled formulas outside the active target set must remain exactly as supplied by the input.
        for d in declarations.values():
            for row, col in d.owned_envelope.cells():
                cell = source[d.sheet].cell(row, col)
                key = (d.sheet, cell.coordinate)
                if cell.data_type == "f" and key not in affected:
                    expected_formulas[key] = cell.value
        require(actual_formulas == expected_formulas, "Dynamic formula inventory readback failed")
        evidence["formulas"] = sorted((sheet, cell, formula) for (sheet, cell), formula in actual_formulas.items())
        for d in declarations.values():
            for row, col in d.owned_envelope.cells():
                address = output[d.sheet].cell(row, col).coordinate
                key = (d.sheet, address)
                if key in affected:
                    continue
                before, after = source[d.sheet][address], output[d.sheet][address]
                require((before.value, before.data_type, before.number_format, before.style_id)
                        == (after.value, after.data_type, after.number_format, after.style_id),
                        "Unplanned Dynamic Region mutation")
                evidence["unchanged"].append((d.sheet, address))
        evidence_hash = fingerprint(evidence)
        return evidence_hash, evidence
    finally:
        for workbook in (source, output):
            workbook.close()
            if workbook.vba_archive:
                workbook.vba_archive.close()


def render_dynamic(request: DynamicRenderRequest, input_path: Path, output_path: Path, *, plan=None):
    input_path, output_path = Path(input_path), Path(output_path)
    require(output_path.suffix.lower() == ".xlsm" and not output_path.exists(), "Invalid output destination")
    expected = plan_dynamic_render(request, input_path)
    require(plan is None or plan == expected, "Incomplete/tampered Dynamic RenderPlan")
    q = request.qualification
    declarations = {d.region_id: d for d in q.declarations}
    styles = {p.style_policy_id: p for p in q.style_policies}
    formulas = {p.formula_policy_id: p for p in q.formula_policies}
    with TemporaryDirectory(prefix="m7_dynamic_", dir=output_path.parent) as temp:
        staged = Path(temp) / "staged.xlsm"
        w = load_workbook(input_path, keep_vba=True, keep_links=True)
        try:
            for operation in expected.dynamic_operations:
                d, sheet = declarations[operation.region_id], w[declarations[operation.region_id].sheet]
                if operation.operation_type == "RESIZE_TABLE":
                    table = sheet.tables[d.table_id]
                    existing = {column.name: column for column in table.tableColumns}
                    desired = d.field_columns[:operation.requested_columns]
                    table.tableColumns = [existing.get(field.field_id,
                        TableColumn(id=index + 1, name=field.field_id))
                        for index, field in enumerate(desired)]
                    for index, column in enumerate(table.tableColumns, 1):
                        column.id = index
                    table.ref = (f"{get_column_letter(operation.after_bounds.left)}{operation.after_bounds.top - 1}:"
                                 f"{get_column_letter(operation.after_bounds.right)}{operation.after_bounds.bottom}")
                elif operation.operation_type == "CLEAR_OWNED_STALE":
                    for address in operation.affected_cells:
                        sheet[address.split("!", 1)[1]].value = None
                elif operation.operation_type == "PROPAGATE_STYLE":
                    policy = styles[d.style_policy_ref]
                    source = sheet[policy.template_cell.split("!", 1)[1]]
                    for address in operation.affected_cells:
                        _apply_style(source, sheet[address.split("!", 1)[1]], policy)
                    if "row_height" in policy.owned_attributes:
                        for row in range(operation.after_bounds.top, operation.after_bounds.bottom + 1):
                            sheet.row_dimensions[row].height = sheet.row_dimensions[source.row].height
                    if "column_width" in policy.owned_attributes:
                        for col in range(operation.after_bounds.left, operation.after_bounds.right + 1):
                            sheet.column_dimensions[get_column_letter(col)].width = sheet.column_dimensions[get_column_letter(source.column)].width
                elif operation.operation_type == "PROPAGATE_PRESENTATION_FORMULA":
                    policy = formulas[d.formula_policy_ref]
                    origin = policy.source_formula_cells[0].split("!", 1)[1]
                    formula = sheet[origin].value
                    for address in operation.affected_cells:
                        target = address.split("!", 1)[1]
                        sheet[target] = Translator(formula, origin=origin).translate_formula(target)
            for write in expected.writes:
                cell = w[write.sheet][write.cell]
                cell.value = literal(write.value) if write.kind == "literal" else write.value
                cell.number_format = write.number_format
            w.save(staged)
        finally:
            w.close()
            if w.vba_archive:
                w.vba_archive.close()
        require(_vba(staged) == expected.vba_sha256, "VBA preservation failure")
        require(_protected(staged, q.declarations) == expected.protected_parts_sha256,
                "Protected surface mutation")
        actual = plan_dynamic_render(request if request.lineage is None else request, input_path)
        require(actual == expected, "Dynamic plan changed during render")
        out_state = _logical_state(staged, q.declarations)
        require(all(out_state[d.region_id] for d in q.declarations), "Incomplete logical output")
        validation_sha, validation = _validate_dynamic_output(input_path, staged, request, expected)
        shutil.copyfile(staged, output_path)
    return {"qa": "PASS", "plan_sha256": expected.plan_sha256,
            "output_sha256": sha(output_path), "logical_output_sha256": fingerprint(_logical_state(output_path, q.declarations)),
            "vba_before": expected.vba_sha256, "vba_after": _vba(output_path),
            "protected_parts_before": expected.protected_parts_sha256,
            "output_validation_sha256": validation_sha, "output_validation": validation,
            "region_bounds": tuple((d.region_id, next(o.after_bounds for o in expected.dynamic_operations
                                    if o.region_id == d.region_id)) for d in q.declarations)}


def lineage_from(plan, evidence):
    return DynamicLineage(plan.qualified_master_sha256, evidence["output_sha256"],
                          evidence["region_bounds"], plan.plan_sha256)
