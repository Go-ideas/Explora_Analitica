from __future__ import annotations

from dataclasses import dataclass, field
from io import BytesIO
import json
from pathlib import Path

import numpy as np
import pandas as pd
from scipy.stats import binomtest, t, ttest_rel
from openpyxl import load_workbook

from src.database.db_reader import read_table, list_tables
from src.export.excel_exporter import export_report_table_to_excel
from src.reporter.filters import apply_multiple_filters
from src.utils.text_utils import normalize_text


@dataclass(frozen=True)
class PairedCondition:
    question_id: str
    filters: dict[str, list[object]] = field(default_factory=dict)


@dataclass
class PairedComparison:
    table: pd.DataFrame
    config: dict
    notes: list[str]

    def excel_bytes(self) -> bytes:
        output = BytesIO()
        export_report_table_to_excel(self.table, self.config, output, notes=self.notes)
        workbook = load_workbook(BytesIO(output.getvalue()))
        sheet = workbook['Tabla']
        for row in sheet:
            for cell in row:
                if cell.value == 'p bilateral':
                    for index in range(cell.row + 1, sheet.max_row + 1):
                        sheet.cell(index, cell.column).number_format = '0.000000E+00'
        final = BytesIO()
        workbook.save(final)
        return final.getvalue()


def _describe(condition: PairedCondition) -> str:
    return condition.question_id + ' | ' + json.dumps(condition.filters, ensure_ascii=False, sort_keys=True, default=str)


def compare_paired(
    db_path: Path, left: PairedCondition, right: PairedCondition,
    *, metric: str = 'Media', confidence: float = .95, min_pairs: int = 30,
    duplicate_policy: str = 'reject', option_code: object = None,
) -> PairedComparison:
    """One explicit, unweighted contrast using only complete person pairs.

    Means use a paired t-test; binary outcomes use exact McNemar. Repeated
    observations within a side require an explicit person-mean policy for means.
    Neither row order nor an arbitrary first observation defines a pair.
    """
    if not np.isfinite(confidence) or not 0 < confidence < 1 or not isinstance(min_pairs, int) or min_pairs < 2:
        raise ValueError('Confianza o mínimo de pares inválidos.')
    if metric not in {'Media', 'Top2Box', 'BottomBox', 'Selección RM'}:
        raise ValueError('Métrica pareada no soportada.')
    if duplicate_policy not in {'reject', 'person_mean'}:
        raise ValueError('Política de repeticiones desconocida.')
    if metric != 'Media' and duplicate_policy != 'reject':
        raise ValueError('Las proporciones pareadas requieren una observación por persona y condición.')
    if left == right:
        raise ValueError('Seleccione dos condiciones diferentes.')
    respondents = read_table(db_path, 'respondentes')
    if not {'id_persona','id_original'}.issubset(respondents.columns) or respondents.id_persona.isna().any():
        raise ValueError('La base necesita identidad de persona; regenérela con el Datamap.')
    identity = respondents.id_persona.astype(str).str.strip()
    if identity.eq('').any() or not identity.eq(respondents.id_original.astype(str).str.strip()).all():
        raise ValueError('Se requiere un identificador de persona válido del estudio; no se usan identificadores generados por fila para inferir pares.')
    questions = read_table(db_path, 'preguntas').set_index('pregunta_id')
    for condition in (left, right):
        if condition.question_id not in questions.index or not bool(int(questions.at[condition.question_id, 'usar_en_dashboard'])):
            raise ValueError('Una pregunta no existe o está excluida del análisis.')
        kind = normalize_text(questions.loc[condition.question_id].get('tipo_pregunta'))
        is_rm = kind in {'rm', 'rm_codificada'}
        is_scale = any(token in kind for token in ('escala', 'nps', 'numerica'))
        if (metric == 'Selección RM') != is_rm or (metric != 'Selección RM' and not is_scale):
            raise ValueError('La métrica no corresponde al tipo de pregunta de una condición.')
        for variable, values in condition.filters.items():
            if variable not in respondents or not isinstance(values, list) or not values:
                raise ValueError(f'Filtro ausente o inválido: {variable}')
    if left.question_id != right.question_id:
        a = questions.loc[left.question_id].get('pregunta_padre')
        b = questions.loc[right.question_id].get('pregunta_padre')
        if pd.isna(a) or not str(a).strip() or a != b:
            raise ValueError('Las preguntas deben compartir una pregunta padre declarada para comparar ítems repetidos.')
    if metric == 'Selección RM':
        numeric_code = pd.to_numeric(option_code, errors='coerce')
        if pd.isna(numeric_code) or not np.isfinite(numeric_code) or numeric_code == 0:
            raise ValueError('Seleccione un código RM numérico válido; cero significa no seleccionado.')
    table_name = 'multirrespuesta_long' if metric == 'Selección RM' else 'escalas_long'
    source = read_table(db_path, table_name)
    universes = None
    if metric == 'Selección RM':
        if 'bases_multirrespuesta' not in list_tables(db_path):
            raise ValueError('La RM requiere bases observadas explícitas; regenere la base con el formato declarado.')
        universes = read_table(db_path, 'bases_multirrespuesta')
    options = read_table(db_path, 'opciones')
    if left.question_id != right.question_id:
        definitions = []
        for condition in (left, right):
            declared = options[options.pregunta_id.eq(condition.question_id)]
            declared = declared[declared.codigo.notna()]
            definitions.append(set(zip(declared.codigo.map(_code), declared.label.map(normalize_text))))
        if definitions[0] != definitions[1]:
            raise ValueError('Las etiquetas o códigos de respuesta no coinciden entre los ítems; valide su equivalencia antes de comparar.')
    frames = []
    bounds = []
    selected_rows = []
    observation_counts = []
    repeat_counts = []
    for condition in (left, right):
        people = apply_multiple_filters(respondents, condition.filters)
        rows = source[source.pregunta_id.eq(condition.question_id)].copy()
        if metric == 'Selección RM':
            allowed = options[options.pregunta_id.eq(condition.question_id)]
            if not allowed.codigo.astype(str).map(_code).eq(_code(option_code)).any():
                raise ValueError('La opción RM no está declarada en una de las condiciones.')
            rows = universes[universes.pregunta_id.eq(condition.question_id)].copy()
            selected = source[source.pregunta_id.eq(condition.question_id)]
            if rows.empty and not selected.empty:
                raise ValueError('Esta RM no tiene una base observada explícita; revise el formato y regenere la base.')
            selected = set(selected.loc[selected.codigo_respuesta.map(_code).eq(_code(option_code)), 'id_respondente'])
            rows['value'] = rows.id_respondente.isin(selected).astype(int)
        else:
            if rows.variable.nunique() > 1:
                raise ValueError('Seleccione un ítem escalar separado; no se promedian distintas variables dentro de una condición.')
            column = {'Media':'valor', 'Top2Box':'top2box', 'BottomBox':'bottombox'}[metric]
            if metric != 'Media' and not rows.empty and rows[column].isna().all():
                raise ValueError('La caja requiere límites de escala definidos en ambas condiciones.')
            rows['value'] = pd.to_numeric(rows[column], errors='coerce')
            definition = rows[['escala_min','escala_max']].drop_duplicates().dropna()
            bounds.append(tuple(map(tuple, definition.to_numpy())))
        rows = rows.merge(people[['id_respondente','id_persona']], on='id_respondente', how='inner', validate='many_to_one')
        rows = rows[rows.value.notna()]
        if not np.isfinite(rows.value.to_numpy(dtype=float)).all():
            raise ValueError('Hay valores no finitos en la comparación.')
        if metric != 'Media' and not rows.value.isin([0,1]).all():
            raise ValueError('La comparación de proporciones requiere estados binarios definidos.')
        selected_rows.append(set(rows.id_respondente))
        observation_counts.append(len(rows))
        repeat_counts.append(int(rows.groupby('id_persona').size().gt(1).sum()))
        if rows.id_persona.duplicated().any() and duplicate_policy == 'reject':
            raise ValueError('Hay varias evaluaciones por persona en una condición. Filtre una sola o elija explícitamente la media por persona para comparar medias.')
        frames.append(rows.groupby('id_persona').value.mean())
    if left.question_id == right.question_id and selected_rows[0] & selected_rows[1]:
        raise ValueError('Las condiciones se solapan: una misma evaluación aparece en ambos lados.')
    if metric != 'Selección RM' and (bounds[0] != bounds[1] or len(bounds[0]) > 1):
        raise ValueError('Las condiciones tienen escalas incompatibles o mezcladas.')
    paired = pd.concat([frames[0].rename('a'), frames[1].rename('b')], axis=1).dropna()
    n = len(paired)
    p_value = low = high = None
    a_mean = float(paired.a.mean()) if n else None
    b_mean = float(paired.b.mean()) if n else None
    effect = a_mean - b_mean if n else None
    if n and not np.isfinite([a_mean, b_mean, effect]).all():
        raise ValueError('No se puede calcular una diferencia finita con los valores de estas condiciones.')
    method = 't pareada bilateral' if metric == 'Media' else 'McNemar exacta bilateral'
    status = 'Base insuficiente'
    discordant_ab = discordant_ba = None
    if n >= min_pairs:
        if metric == 'Media':
            differences = paired.a - paired.b
            sd = differences.std(ddof=1)
            if not np.isfinite(sd):
                raise ValueError('La variabilidad de las diferencias no permite calcular el contraste.')
            if sd == 0:
                p_value = 1.0 if effect == 0 else 0.0
                low = high = effect
            else:
                p_value = float(ttest_rel(paired.a, paired.b).pvalue)
                if not np.isfinite(p_value):
                    raise ValueError('La prueba no produjo un p-valor válido.')
                margin = float(t.ppf((1+confidence)/2, n-1) * sd / np.sqrt(n))
                low, high = effect-margin, effect+margin
        else:
            discordant_ab = int((paired.a.eq(1) & paired.b.eq(0)).sum())
            discordant_ba = int((paired.a.eq(0) & paired.b.eq(1)).sum())
            discordant = discordant_ab + discordant_ba
            p_value = float(binomtest(discordant_ab, discordant, .5).pvalue) if discordant else 1.0
        status = 'Diferencia significativa' if p_value < 1-confidence else 'Sin diferencia significativa'
    result = pd.DataFrame([{
        'Métrica':metric, 'Personas A disponibles':len(frames[0]), 'Personas B disponibles':len(frames[1]),
        'Evaluaciones A disponibles':observation_counts[0], 'Evaluaciones B disponibles':observation_counts[1],
        'Personas repetidas A':repeat_counts[0], 'Personas repetidas B':repeat_counts[1],
        'Pares completos':n, 'Sin pareja A':len(frames[0])-n, 'Sin pareja B':len(frames[1])-n,
        'Valor A (pares)':a_mean, 'Valor B (pares)':b_mean, 'Diferencia A−B':effect,
        'IC inferior':low, 'IC superior':high, 'p bilateral':p_value,
        'A sí/B no':discordant_ab, 'A no/B sí':discordant_ba, 'Resultado':status,
    }])
    notes = [
        'Solo se comparan personas con datos válidos en ambas condiciones; los valores A/B se calculan sobre esos mismos pares.',
        'Contraste bilateral único, sin ponderación y sin ajuste por comparaciones múltiples. No controla orden, arrastre ni otros efectos del diseño.',
        'Una diferencia significativa no demuestra causalidad ni que un código alto sea comercialmente mejor.',
        'La política media por persona promedia las evaluaciones de cada lado antes del test; no elige una fila arbitraria.' if duplicate_policy == 'person_mean' else 'Se exige una sola observación por persona y condición.',
        'La t pareada supone independencia entre personas y un modelo adecuado para sus diferencias.' if metric == 'Media' else 'McNemar usa los pares discordantes. El intervalo de diferencia de proporciones no se estima en esta salida.',
    ]
    return PairedComparison(result, {
        'Pregunta':f'Comparación pareada · {metric}', 'Condición A':_describe(left), 'Condición B':_describe(right),
        'Método':method, 'Confianza':confidence, 'Mínimo de pares':min_pairs,
        'Repeticiones':duplicate_policy, 'Código RM':str(option_code) if option_code is not None else '',
        'Base válida':n, 'Unidad de base':'Pares de personas',
    }, notes)


def _code(value: object) -> str:
    try:
        return str(float(value))
    except (ValueError, TypeError):
        return str(value)
