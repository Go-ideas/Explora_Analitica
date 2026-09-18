"""Gate 36 test infrastructure only; not a CanonicalResult renderer."""
from __future__ import annotations

import hashlib
from io import BytesIO
import json
import posixpath
import shutil
from pathlib import Path
from zipfile import BadZipFile, ZipFile
from xml.etree import ElementTree as ET

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.table import Table
from openpyxl.workbook.defined_name import DefinedName

S = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
R = "http://schemas.openxmlformats.org/package/2006/relationships"
MACRO = "application/vnd.ms-excel.sheet.macroEnabled.main+xml"
VBA_REL = "http://schemas.microsoft.com/office/2006/relationships/vbaProject"
ROLES = ("presentation", "provenance", "results", "bases", "slices",
         "significance", "qa", "cell_map", "canonical_payload",
         "configuration", "navigation")


def digest(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def fingerprint(value) -> str:
    return digest(json.dumps(value, sort_keys=True, separators=(",", ":")).encode())


def xml_hash(data: bytes) -> str:
    return digest(ET.canonicalize(data).encode())


def validate_vba_relationships(rels) -> None:
    vba_rels = [e for e in rels if e.get("Type") == VBA_REL]
    if len(vba_rels) != 1 or vba_rels[0].get("TargetMode") == "External":
        raise ValueError("VBA relationship missing or invalid")
    target = vba_rels[0].get("Target", "")
    resolved = target.lstrip("/") if target.startswith("/") else posixpath.normpath("xl/" + target)
    if resolved != "xl/vbaProject.bin":
        raise ValueError("VBA relationship target invalid")


def inspect_package(path: Path, *, require_vba: bool = True,
                    ignored_cell: str | None = None) -> dict:
    if require_vba and path.suffix.lower() != ".xlsm":
        raise ValueError("VBA source requires .xlsm")
    try:
        with ZipFile(path) as z:
            if z.testzip() is not None or len(z.namelist()) != len(set(z.namelist())):
                raise ValueError("Corrupt or duplicate package parts")
            contents = ET.fromstring(z.read("[Content_Types].xml"))
            types = {e.get("PartName"): e.get("ContentType") for e in contents}
            rels = ET.fromstring(z.read("xl/_rels/workbook.xml.rels"))
            vba_rels = [e for e in rels if e.get("Type") == VBA_REL]
            if require_vba:
                if types.get("/xl/workbook.xml") != MACRO:
                    raise ValueError("Macro-enabled content type missing")
                payload = z.read("xl/vbaProject.bin")
                if not payload.startswith(bytes.fromhex("D0CF11E0A1B11AE1")):
                    raise ValueError("VBA payload is not an OLE compound container")
                validate_vba_relationships(rels)
            book = ET.fromstring(z.read("xl/workbook.xml"))
            sheets = [(digest(e.get("name", "").encode()), e.get("sheetId"),
                       e.get("state", "visible")) for e in book.findall(f"{{{S}}}sheets/{{{S}}}sheet")]
            names = [(digest(e.get("name", "").encode()), fingerprint(e.attrib),
                      digest((e.text or "").encode())) for e in book.findall(f"{{{S}}}definedNames/{{{S}}}definedName")]
            parts = {}
            formulas = []
            for name in sorted(z.namelist()):
                if z.getinfo(name).file_size > 64 * 1024 * 1024:
                    hasher = hashlib.sha256()
                    with z.open(name) as stream:
                        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
                            hasher.update(chunk)
                    checksum = hasher.hexdigest()
                    parts[name] = {"sha256": checksum, "semantic": checksum}
                    continue
                data = z.read(name)
                raw_hash = digest(data)
                if name == "docProps/core.xml":
                    core = ET.fromstring(data)
                    for field in list(core):
                        if field.tag == "{http://purl.org/dc/terms/}modified":
                            core.remove(field)
                    data = ET.tostring(core)
                if name.startswith("xl/worksheets/sheet") and name.endswith(".xml"):
                    if name != "xl/worksheets/sheet1.xml":
                        for _, cell in ET.iterparse(BytesIO(data), events=("end",)):
                            if cell.tag == f"{{{S}}}c":
                                formula = cell.find(f"{{{S}}}f")
                                if formula is not None:
                                    formulas.append((name, cell.get("r"), fingerprint(formula.attrib), digest((formula.text or "").encode())))
                                cell.clear()
                            elif cell.tag == f"{{{S}}}row":
                                cell.clear()
                        parts[name] = {"sha256": raw_hash, "semantic": raw_hash}
                        continue
                    tree = ET.fromstring(data)
                    for row in tree.findall(f"{{{S}}}sheetData/{{{S}}}row"):
                        if ignored_cell and row.get("r") == "".join(c for c in ignored_cell if c.isdigit()) and name == "xl/worksheets/sheet1.xml":
                            row.attrib.pop("spans", None)
                        removed = False
                        for cell in list(row):
                            if name == "xl/worksheets/sheet1.xml" and cell.get("r") == ignored_cell:
                                row.remove(cell)
                                removed = True
                                continue
                            formula = cell.find(f"{{{S}}}f")
                            if formula is not None:
                                formulas.append((name, cell.get("r"), fingerprint(formula.attrib), digest((formula.text or "").encode())))
                        if removed and not list(row):
                            tree.find(f"{{{S}}}sheetData").remove(row)
                    # The designated technical cell can legitimately enlarge dimension.
                    if ignored_cell and name == "xl/worksheets/sheet1.xml":
                        dim = tree.find(f"{{{S}}}dimension")
                        if dim is not None:
                            tree.remove(dim)
                    data = ET.tostring(tree)
                parts[name] = {"sha256": raw_hash,
                               "semantic": xml_hash(data) if name.endswith((".xml", ".rels")) and len(data) <= 1024 * 1024 else digest(data)}
            groups = {key: [n for n in parts if match(n)] for key, match in {
                "tables": lambda n: n.startswith("xl/tables/") and n.endswith(".xml"),
                "charts": lambda n: n.startswith("xl/charts/") and n.endswith(".xml"),
                "pivots": lambda n: "pivot" in n.lower(),
                "slicers": lambda n: "slicer" in n.lower(),
                "connections": lambda n: n == "xl/connections.xml",
            }.items()}
            return {"whole_sha256": digest(path.read_bytes()), "size": path.stat().st_size,
                    "vba_sha256": digest(z.read("xl/vbaProject.bin")) if "xl/vbaProject.bin" in parts else None,
                    "macro_type": types.get("/xl/workbook.xml") == MACRO,
                    "vba_relationships": fingerprint([e.attrib for e in vba_rels]),
                    "sheets": sheets, "defined_names": names,
                    "formulas": fingerprint(sorted(formulas)), "formula_count": len(formulas),
                    "groups": groups, "parts": parts}
    except (BadZipFile, KeyError, ET.ParseError, UnicodeError) as exc:
        raise ValueError("Invalid workbook package") from exc


def compare(before: dict, after: dict) -> dict:
    a, b = before["parts"], after["parts"]
    return {"vba_exact": before["vba_sha256"] == after["vba_sha256"],
            "sheets": before["sheets"] == after["sheets"],
            "names": before["defined_names"] == after["defined_names"],
            "formulas": before["formulas"] == after["formulas"],
            "vba_relationships": before["vba_relationships"] == after["vba_relationships"],
            "missing_parts": sorted(set(a) - set(b)),
            "added_parts": sorted(set(b) - set(a)),
            "changed_semantic_parts": [n for n in sorted(set(a) & set(b)) if a[n]["semantic"] != b[n]["semantic"]]}


def guarded_destination(source: Path, destination: Path, *, preserve_vba: bool) -> None:
    if not preserve_vba or destination.suffix.lower() != ".xlsm":
        raise ValueError("Preservation mode and .xlsm destination required")
    if source.resolve() == destination.resolve():
        raise ValueError("Source mutation prohibited")
    inspect_package(source)


def create_generic(path: Path) -> None:
    w = Workbook()
    w.remove(w.active)
    for role in ROLES:
        sheet = w.create_sheet(role)
        sheet.sheet_state = "hidden" if role in ("results", "bases", "slices", "significance", "qa", "cell_map", "canonical_payload") else "visible"
        sheet.protection.sheet = True
        if role not in ("presentation", "configuration"):
            sheet.append(["record_id", "value"])
            sheet.append([role + "_RECORD_1", "GENERIC_ONLY"])
            sheet.add_table(Table(displayName="Certification_" + role, ref="A1:B2"))
            w.defined_names.add(DefinedName("Slot_" + role, attr_text=f"'{role}'!$B$2"))
    s = w["presentation"]
    s.append(["record_id", "value"])
    s.append(["GENERIC_VALUE_1", 0.125])
    s.add_table(Table(displayName="CertificationResults", ref="A1:B2"))
    s["B2"].number_format = "0.0%"
    s["D1"] = "MASTER_OWNED"
    s["D2"] = "VBA_OWNED_CONCEPT_ONLY_NO_MACRO"
    s["D3"] = "USER_EDITABLE"
    from openpyxl.styles import Protection
    s["D3"].protection = Protection(locked=False)
    s["E1"] = "=1+1"  # Certification-only authorized formula, not analytics.
    w.defined_names.add(DefinedName("RendererSlot", attr_text="'presentation'!$B$2"))
    w.defined_names.add(DefinedName("LiteralSlot", attr_text="'presentation'!$F$1"))
    w.defined_names.add(DefinedName("FormulaSlot", attr_text="'presentation'!$E$1"))
    w["provenance"]["D1"] = "M7_MASTER_INTERFACE_CONTRACT_V1"
    w["provenance"]["D2"] = "GENERIC_CERTIFICATION_ONLY"
    registry = w["configuration"]
    registry.append(["role_id", "sheet", "owner", "slot_id", "table_binding", "capacity", "policy"])
    for role in ROLES:
        registry.append([role, role, "MASTER" if role in ("configuration", "navigation") else "RENDERER",
                         "RendererSlot" if role == "presentation" else "Slot_" + role,
                         "CertificationResults" if role == "presentation" else ("MasterRoleRegistry" if role == "configuration" else "Certification_" + role),
                         1, "preserve" if role in ("configuration", "navigation") else "replace_declared_only"])
    registry.add_table(Table(displayName="MasterRoleRegistry", ref=f"A1:G{len(ROLES) + 1}"))
    w.defined_names.add(DefinedName("Slot_configuration", attr_text="'configuration'!$A$2"))
    w.save(path)
    w.close()


def check_interface(path: Path) -> None:
    with ZipFile(path) as package:
        book = ET.fromstring(package.read("xl/workbook.xml"))
        names = {s.get("name") for s in book.findall(f"{{{S}}}sheets/{{{S}}}sheet")}
        if not set(ROLES).issubset(names):
            raise ValueError("Required semantic role missing")
        if any(p.file_size > 64 * 1024 * 1024 and p.filename.endswith(".xml") for p in package.infolist()):
            raise ValueError("Generic interface inspection resource bound exceeded")
    w = load_workbook(path)
    try:
        if any(role not in w.sheetnames for role in ROLES):
            raise ValueError("Required semantic role missing")
        if "RendererSlot" not in w.defined_names:
            raise ValueError("Mandatory defined name missing")
        if "CertificationResults" not in w["presentation"].tables:
            raise ValueError("Mandatory table missing")
        if "MasterRoleRegistry" not in w["configuration"].tables:
            raise ValueError("Semantic role registry missing")
        if w["provenance"]["D1"].value != "M7_MASTER_INTERFACE_CONTRACT_V1":
            raise ValueError("Interface version mismatch")
        rows = list(w["configuration"].iter_rows(min_row=2, max_row=len(ROLES)+1, values_only=True))
        if {row[0] for row in rows} != set(ROLES):
            raise ValueError("Role registry identity mismatch")
        for role, sheet, owner, slot, table, capacity, policy in rows:
            expected_owner = "MASTER" if role in ("configuration", "navigation") else "RENDERER"
            if sheet != role or owner != expected_owner or capacity != 1:
                raise ValueError("Role ownership/capacity mismatch")
            if slot not in w.defined_names or table not in w[sheet].tables:
                raise ValueError("Role binding missing")
            expected_cell = "A2" if role == "configuration" else "B2"
            destinations = list(w.defined_names[slot].destinations)
            if len(destinations) != 1 or destinations[0][0] != sheet or destinations[0][1].replace("$", "") != expected_cell:
                raise ValueError("Role named-range destination mismatch")
    finally:
        w.close()


def literal_write(cell, text: str) -> None:
    cell.value = text
    cell.data_type = "s"


def formula_write(cell, formula: str) -> None:
    if cell.coordinate != "E1" or formula != "=1+1":
        raise ValueError("Formula not allowlisted for certification")
    cell.value = formula


def targeted_write(source: Path, destination: Path, address: str, value,
                   *, kind: str = "literal", require_vba: bool = True,
                   preserve_vba: bool = True, require_interface: bool = True,
                   role: str = "presentation") -> None:
    """Declared-cell certification prototype, not a general workbook engine."""
    from lxml import etree
    from openpyxl.utils.cell import range_boundaries, get_column_letter

    if source.resolve() == destination.resolve():
        raise ValueError("Source mutation prohibited")
    if require_vba:
        guarded_destination(source, destination, preserve_vba=preserve_vba)
    elif source.suffix.lower() != ".xlsx" or destination.suffix.lower() != ".xlsx":
        raise ValueError("Generic fixture requires .xlsx")
    if require_interface:
        check_interface(source)
        if role not in ROLES or role in ("configuration", "navigation"):
            raise ValueError("Master-owned or unknown role")
        allowed = {"B2": "numeric", "F1": "literal", "E1": "formula"} if role == "presentation" else {"B2": kind}
        if role != "presentation" and kind == "formula":
            raise ValueError("Formula not permitted in audit roles")
        if allowed.get(address) != kind:
            raise ValueError("Non-renderer slot or incompatible write")
    elif address != "XFD1" or kind != "literal" or value != "CERTIFICATION_ONLY":
        raise ValueError("External stress operation is restricted")
    if kind == "formula" and (address != "E1" or value != "=1+1"):
        raise ValueError("Formula not allowlisted")
    if kind == "literal" and not isinstance(value, str):
        raise ValueError("Literal text required")
    if kind == "numeric":
        import math
        if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value):
            raise ValueError("Finite typed numeric value required")
        if len(str(value).replace(".", "").replace("-", "").replace("+", "")) > 15:
            raise ValueError("Unsafe numeric precision")
    if kind not in ("literal", "numeric", "formula"):
        raise ValueError("Unknown write kind")
    staged = destination.with_name(destination.name + ".stage")
    if destination.exists() or staged.exists():
        raise ValueError("Destination/stage already exists")
    try:
        with ZipFile(source) as original:
            if any("signature" in name.lower() for name in original.namelist()):
                raise ValueError("Signed package requires separate signature certification")
            edited_part = "xl/worksheets/sheet1.xml"
            if require_interface:
                book = ET.fromstring(original.read("xl/workbook.xml"))
                rels = ET.fromstring(original.read("xl/_rels/workbook.xml.rels"))
                rid = next(e for e in book.findall(f"{{{S}}}sheets/{{{S}}}sheet") if e.get("name") == role).get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
                target = next(e for e in rels if e.get("Id") == rid).get("Target")
                edited_part = target.lstrip("/") if target.startswith("/") else posixpath.normpath("xl/" + target)
            parser = etree.XMLParser(resolve_entities=False, no_network=True)
            tree = etree.fromstring(original.read(edited_part), parser)
            data = tree.find(f"{{{S}}}sheetData")
            row_number = int("".join(c for c in address if c.isdigit()))
            row = next((r for r in data if r.get("r") == str(row_number)), None)
            if row is None:
                row = etree.Element(f"{{{S}}}row", r=str(row_number))
                index = next((i for i, r in enumerate(data) if int(r.get("r")) > row_number), len(data))
                data.insert(index, row)
            cell = next((c for c in row if c.get("r") == address), None)
            if not require_interface and cell is not None:
                raise ValueError("External technical cell is occupied")
            if cell is None:
                cell = etree.SubElement(row, f"{{{S}}}c", r=address)
            for child in list(cell):
                cell.remove(child)
            if kind == "literal":
                cell.set("t", "inlineStr")
                inline = etree.SubElement(cell, f"{{{S}}}is")
                text = etree.SubElement(inline, f"{{{S}}}t")
                text.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
                text.text = value
            elif kind == "numeric":
                cell.set("t", "n")
                etree.SubElement(cell, f"{{{S}}}v").text = repr(value)
            else:
                cell.attrib.pop("t", None)
                etree.SubElement(cell, f"{{{S}}}f").text = value[1:]
            dim = tree.find(f"{{{S}}}dimension")
            if "spans" in row.attrib:
                start, end = map(int, row.get("spans").split(":"))
                row.set("spans", f"{min(start,range_boundaries(address)[0])}:{max(end,range_boundaries(address)[0])}")
            if dim is not None:
                left, top, right, bottom = range_boundaries(dim.get("ref"))
                col = range_boundaries(address)[0]
                dim.set("ref", f"{get_column_letter(min(left,col))}{min(top,row_number)}:{get_column_letter(max(right,col))}{max(bottom,row_number)}")
            edited = etree.tostring(tree, xml_declaration=True, encoding="UTF-8", standalone=True)
            with ZipFile(staged, "w") as output:
                for item in original.infolist():
                    if item.filename == edited_part:
                        output.writestr(item, edited)
                    else:
                        with original.open(item) as incoming, output.open(item, "w") as outgoing:
                            shutil.copyfileobj(incoming, outgoing, 1024 * 1024)
            inspect_package(staged, require_vba=False)
            # Validate macro source/output without relying on the staging extension.
            if require_vba:
                with ZipFile(source) as before, ZipFile(staged) as after:
                    if before.read("xl/vbaProject.bin") != after.read("xl/vbaProject.bin"):
                        raise ValueError("VBA payload changed")
                    for item in before.namelist():
                        if item != edited_part:
                            def part_hash(archive, name):
                                h = hashlib.sha256()
                                with archive.open(name) as part:
                                    for chunk in iter(lambda: part.read(1024 * 1024), b""):
                                        h.update(chunk)
                                return h.digest()
                            if part_hash(before, item) != part_hash(after, item):
                                raise ValueError("Untouched package part changed")
            staged.replace(destination)
    finally:
        staged.unlink(missing_ok=True)
