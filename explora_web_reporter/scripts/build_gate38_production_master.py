"""Build the generic Gate 38 Production Master candidate from public test infrastructure."""
from __future__ import annotations

import argparse
from pathlib import Path
import sys

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table
from openpyxl.workbook.defined_name import DefinedName

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from src.excel_renderer.renderer import ROLES, audit_columns


def build(seed: Path, output: Path) -> None:
    if output.exists():
        raise ValueError("Production Master destination must not exist")
    workbook = load_workbook(seed, keep_vba=True, keep_links=True)
    try:
        role_tables = {}
        for role in ROLES:
            sheet = workbook[role]
            if role in ("presentation", "configuration", "navigation"):
                table = next(iter(sheet.tables.values()))
                role_tables[role] = (table.name, 11 if role == "configuration" else 1)
                continue
            for name in list(sheet.tables):
                del sheet.tables[name]
            columns = audit_columns(role)
            capacity = 256
            for column, name in enumerate(columns, 1):
                sheet.cell(1, column, name)
            for row in range(2, capacity + 2):
                for column in range(1, len(columns) + 1):
                    sheet.cell(row, column, "MASTER_RESERVED")
            name = "Production_" + role
            sheet.add_table(Table(displayName=name,
                                  ref=f"A1:{get_column_letter(len(columns))}{capacity + 1}"))
            role_tables[role] = (name, capacity)
        sheet = workbook["presentation"]
        sheet["A2"] = "EXPLORA_PRODUCTION_MASTER_V1"
        sheet["B2"] = 0
        sheet["B2"].number_format = "0"
        sheet["B3"] = 0
        sheet["B3"].number_format = "0"
        sheet["C1"] = "RESULT_LABEL"
        sheet["C2"] = "BASE_LABEL"
        sheet["F1"] = "PROJECT_LABEL"
        sheet["G1"] = "RESULT_STATUS"
        sheet["G2"] = "BASE_STATUS"
        sheet["H1"] = "LABEL_STATUS"
        names = {
            "ProductionResultSlot": "'presentation'!$B$2",
            "ProductionResultLabel": "'presentation'!$C$1",
            "ProductionResultStatus": "'presentation'!$G$1",
            "ProductionBaseSlot": "'presentation'!$B$3",
            "ProductionBaseLabel": "'presentation'!$C$2",
            "ProductionBaseStatus": "'presentation'!$G$2",
            "ProductionProjectLabel": "'presentation'!$F$1",
            "ProductionProjectLabelStatus": "'presentation'!$H$1",
        }
        for name, target in names.items():
            if name in workbook.defined_names:
                del workbook.defined_names[name]
            workbook.defined_names.add(DefinedName(name, attr_text=target))
        workbook.properties.title = "EXPLORA Production Master V1"
        workbook.properties.subject = "Generic contract-driven canonical result presentation"
        workbook.properties.keywords = "EXPLORA,Production Master,M7"
        registry = workbook["configuration"]
        for index, role in enumerate(ROLES, 2):
            table, capacity = role_tables[role]
            registry.cell(index, 5, table)
            registry.cell(index, 6, capacity)
        output.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(output)
    finally:
        workbook.close()
        workbook.vba_archive.close()


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--seed", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    build(args.seed, args.output)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
