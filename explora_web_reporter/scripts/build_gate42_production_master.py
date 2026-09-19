"""Derive the versioned Gate 42 Production Master candidate from accepted 1.1.0."""
from copy import copy
import hashlib
import json
from pathlib import Path
from datetime import datetime
import re
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.workbook.defined_name import DefinedName

ROOT = Path(__file__).resolve().parents[1]
MASTER_DIR = ROOT / "production_masters"
BASELINE = MASTER_DIR / "EXPLORA_PRODUCTION_MASTER_V1.xlsm"
BASELINE_MANIFEST = BASELINE.with_suffix(".qualification.json")
CANDIDATE = MASTER_DIR / "EXPLORA_PRODUCTION_MASTER_V1_2.xlsm"
CANDIDATE_MANIFEST = CANDIDATE.with_suffix(".qualification.json")
DYNAMIC_MANIFEST = CANDIDATE.with_suffix(".dynamic.qualification.json")
BASELINE_SHA = "a8313adf706016b30b8837a7bc42cb80c7ccc281f2663616a7f22b8daa8d10dd"
VBA_SHA = "0f879b60ed12315085bc722c3f59163f86ce24609e3ba6ad379069c44779b758"


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def normalize_package(path):
    staged = path.with_suffix(".normalized.xlsm")
    with ZipFile(path) as source, ZipFile(staged, "w") as target:
        for name in sorted(source.namelist()):
            original = source.getinfo(name)
            info = ZipInfo(name, (2026, 9, 19, 0, 0, 0))
            info.compress_type = original.compress_type if original.compress_type else ZIP_DEFLATED
            info.external_attr = original.external_attr
            payload = source.read(name)
            if name == "docProps/core.xml":
                payload = re.sub(rb"<dcterms:modified[^>]*>.*?</dcterms:modified>",
                    (b'<dcterms:modified xmlns:dcterms="http://purl.org/dc/terms/" '
                     b'xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance" '
                     b'xsi:type="dcterms:W3CDTF">2026-09-19T00:00:00Z</dcterms:modified>'), payload)
            target.writestr(info, payload)
    staged.replace(path)


def build():
    if sha(BASELINE) != BASELINE_SHA:
        raise RuntimeError("Accepted Production Master baseline mismatch")
    workbook = load_workbook(BASELINE, keep_vba=True)
    try:
        workbook.properties.modified = datetime(2026, 9, 19)
        sheet = workbook["presentation"]
        fill = PatternFill("solid", fgColor="FFD9EAF7")
        border = Border(bottom=Side(style="thin", color="FF7F8C8D"))
        for row in range(2, 10):
            for col in range(10, 13):
                cell = sheet.cell(row, col)
                cell.value = None
                cell.font = Font(name="Calibri", size=11, color="FF1F1F1F")
                cell.fill = copy(fill)
                cell.border = copy(border)
                cell.alignment = Alignment(horizontal="center")
                cell.protection = Protection(locked=True)
                cell.number_format = "0.0%" if col == 10 else "General"
        provenance = sheet["N2"]
        provenance.value = None
        provenance.font = Font(name="Calibri", size=10, color="FF404040")
        provenance.fill = PatternFill("solid", fgColor="FFE2F0D9")
        provenance.protection = Protection(locked=True)
        workbook.defined_names.add(DefinedName("ProductionDynamicResultAnchor", attr_text="'presentation'!$J$2"))
        workbook.defined_names.add(DefinedName("ProductionDynamicProvenanceAnchor", attr_text="'presentation'!$N$2"))
        workbook.save(CANDIDATE)
    finally:
        workbook.close()
        if workbook.vba_archive:
            workbook.vba_archive.close()
    normalize_package(CANDIDATE)
    candidate_sha = sha(CANDIDATE)
    fixed = json.loads(BASELINE_MANIFEST.read_text(encoding="utf-8"))
    fixed["master_version"] = "1.2.0"
    fixed["filename"] = CANDIDATE.name
    fixed["master_sha256"] = candidate_sha
    fixed["supported_capabilities"].append("bounded Dynamic Region coexistence")
    CANDIDATE_MANIFEST.write_text(json.dumps(fixed, indent=2) + "\n", encoding="ascii")
    dynamic = {
        "schema_version": "M7_PRODUCTION_DYNAMIC_REGION_QUALIFICATION_V1",
        "master_id": "EXPLORA_PRODUCTION_MASTER_V1",
        "master_version": "1.2.0",
        "filename": CANDIDATE.name,
        "master_sha256": candidate_sha,
        "baseline_version": "1.1.0",
        "baseline_filename": BASELINE.name,
        "baseline_sha256": BASELINE_SHA,
        "vba_project_sha256": VBA_SHA,
        "region_contract_version": "M7_DYNAMIC_REGION_CONTRACT_V1",
        "renderer_build": "M7_DYNAMIC_REGION_RUNTIME_V1",
        "qualification_status": "CANDIDATE",
        "declarations": [{
            "region_id": "PRODUCTION_DYNAMIC_RESULTS_V1", "role_id": "presentation",
            "worksheet_role": "presentation", "sheet": "presentation", "region_kind": "RANGE",
            "table_id": None, "anchor_name": "ProductionDynamicResultAnchor",
            "growth_dimensions": "ROWS_AND_COLUMNS", "min_rows": 1, "min_columns": 1,
            "max_rows": 8, "max_columns": 3,
            "owned_envelope": {"left": 10, "top": 2, "right": 12, "bottom": 9},
            "field_columns": [
                {"field_id": "estimate", "source_field": "estimate", "column_offset": 0,
                 "display_profile_ref": "PROPORTION", "storage_mode": "exact_numeric"},
                {"field_id": "status", "source_field": "value_status", "column_offset": 1,
                 "display_profile_ref": "LITERAL", "storage_mode": "exact_text"},
                {"field_id": "significance", "source_field": "__SIGNIFICANCE_TOKEN__", "column_offset": 2,
                 "display_profile_ref": "SIGNIFICANCE_TOKEN", "storage_mode": "exact_text"}],
            "style_policy_ref": "PRODUCTION_DYNAMIC_RESULTS_STYLE_V1", "formula_policy_ref": "NONE"
        }, {
            "region_id": "PRODUCTION_DYNAMIC_PROVENANCE_V1", "role_id": "provenance",
            "worksheet_role": "presentation", "sheet": "presentation", "region_kind": "RANGE",
            "table_id": None, "anchor_name": "ProductionDynamicProvenanceAnchor",
            "growth_dimensions": "ROWS", "min_rows": 1, "min_columns": 1,
            "max_rows": 1, "max_columns": 1,
            "owned_envelope": {"left": 14, "top": 2, "right": 14, "bottom": 2},
            "field_columns": [{"field_id": "result_run_id", "source_field": "result_run_id", "column_offset": 0,
                               "display_profile_ref": "LITERAL", "storage_mode": "exact_text"}],
            "style_policy_ref": "PRODUCTION_DYNAMIC_PROVENANCE_STYLE_V1", "formula_policy_ref": "NONE"
        }],
        "style_policies": [{
            "style_policy_id": "PRODUCTION_DYNAMIC_RESULTS_STYLE_V1",
            "region_id": "PRODUCTION_DYNAMIC_RESULTS_V1", "template_cell": "presentation!J2",
            "propagation_axis": "ROWS_AND_COLUMNS",
            "owned_attributes": ["font", "fill", "border", "alignment", "protection"],
            "unsupported_attributes": ["row_height", "column_width", "conditional_formatting"],
            "table_style_policy": "NOT_APPLICABLE"
        }, {
            "style_policy_id": "PRODUCTION_DYNAMIC_PROVENANCE_STYLE_V1",
            "region_id": "PRODUCTION_DYNAMIC_PROVENANCE_V1", "template_cell": "presentation!N2",
            "propagation_axis": "ROWS", "owned_attributes": ["font", "fill", "protection"],
            "unsupported_attributes": ["row_height", "column_width", "conditional_formatting"],
            "table_style_policy": "NOT_APPLICABLE"
        }],
        "formula_policies": [],
        "fixed_slots": fixed["writable_slots"],
        "lineage": {"derivation": "DETERMINISTIC_FROM_ACCEPTED_BASELINE", "builder": "scripts/build_gate42_production_master.py"}
    }
    DYNAMIC_MANIFEST.write_text(json.dumps(dynamic, indent=2) + "\n", encoding="ascii")


if __name__ == "__main__":
    build()
